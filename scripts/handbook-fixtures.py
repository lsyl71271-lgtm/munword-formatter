"""Generate deliberately corrupted, bilingual handbook-role fixtures for visual QA."""
import argparse, json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

parser=argparse.ArgumentParser(); parser.add_argument('output'); args=parser.parse_args()
out=Path(args.output); out.mkdir(parents=True,exist_ok=True)
types=['position-paper','working-paper','draft-directive','draft-resolution','friendly-amendment','unfriendly-amendment']
titles=['立场文件','工作文件 2.1','指令草案 1.1','决议草案 1.2','友好修正案 1.2.1','非友好修正案 1.2.1']
english=['Position Paper','Working Paper 4.1','Draft Directive 4.1','Draft Resolution 1.1','Friendly Amendment 1.1.1','Unfriendly Amendment 1.1.1']
manifest=[]
for index,kind in enumerate(types):
    for lang in ['zh','en']:
        lines=[titles[index] if lang=='zh' else english[index]]
        if kind=='position-paper':
            lines += ['委员会：历史安全理事会','议题：纳戈尔诺－卡拉巴赫局势','国家：美利坚合众国','代表：李相霖、张镜如','纳戈尔诺－卡拉巴赫地区是外高加索众多领土争议地区之一，该地区复杂的历史沿革是造成这一领土争议的主要原因。','国际社会对纳卡问题采取的措施是多方面的。从措施的主要推动者来看，可以分为：','一、欧安会明斯克小组采取的措施。','据此，美利坚合众国重申其立场并提出如下倡议：','1、要求各方停止敌对行动。','a) 建立透明的观察机制。'] if lang=='zh' else ['Committee: General Assembly-Legal Committee','Topic: Future Regulations on Soft Drugs in International Law','Country: The Islamic Republic of Afghanistan','Delegate: Li Ying, Nanjing Foreign Languages School','Soft drug, defined and perceived as a drug considered to be non-addictive and less damaging to the health than a hard drug, has been facing debates politically and psychologically.','The extent of global illicit drug use remained stable in the five years up to and including 2010. Despite continued efforts, demand and supply of illicit drugs remained high.','1. Afghanistan strongly suggests all member states to start with the definition and clarification of soft drug;','a) Establish a transparent information-sharing mechanism.']
        else:
            lines += ['委员会：联合国大会','议题：1947年巴勒斯坦问题','起草国：土耳其共和国、叙利亚共和国'] if lang=='zh' else ['Committee: The European Commission','Topic: European Energy Security in the Wake of Russia-Ukraine Crisis','Sponsors: Spain, Germany, France']
            if kind!='working-paper': lines += ['附议国：阿富汗王国、阿根廷共和国、埃及王国、黎巴嫩共和国、沙特阿拉伯王国、伊拉克王国、伊朗国'] if lang=='zh' else ['Signatories: Italy, Croatia, Cyprus, Malta, Slovakia']
            if kind=='working-paper': lines += ['1. 欣见各缔约方对气候变化以及其带来的严重后果的关切，','2. 强调为保障核能安全，需要透明的监督管理机制和严密的防护机制：','(a) 加强各个阶段的安全审查和监管，监督保证技术安全性和质量；','(b) 做好内部管理，保障核废水存储安全；','3. 认可天然气相较传统能源而言具有更小的气候危害性，可作为能源完全脱碳过程中的过渡能源。'] if lang=='zh' else ['Part I: Technological Communication','1. The sponsors agreed that each country should allow countries to share their technology and promote it;','(a) Exchange student programmes;','(b) Foreign professor support;','Part II: Ethics','1. The use of gene engineering must obey common ethics and cannot hurt basic human rights or violate a state’s ethics.']
            elif kind=='draft-directive': lines += ['联合国大会，','1. 要求冲突各方立即停止一切军事行动，安理会将对一切有主动破坏和平行为的国家进行严厉的制裁；','2. 决定派遣观察员对停火状况进行监督并提供报告；','3. 督促冲突各方在停火期间尽快制定出可行的后续计划。'] if lang=='zh' else ['The OPCW,','1. Requests the government of India and Pakistan to fully cooperate with the OPCW in the following:','a. Finding out the modus operandi of the terrorist groups;','b. Allowing full-fledged inspections of the region of Kashmir, where chemical materials are suspected to have been tested;','2. Urges cooperation among security and intelligence related bodies to prevent terrorist attacks that involve chemical agents.']
            elif kind=='draft-resolution': lines += ['联合国大会，','注意到在托管委员会方面达成和平解决争端的建议，','回顾相关原则，','第一条 支持在巴勒斯坦地区成立一个统一的国家；','第二条 决定设置地方自治议会：','（一）地方自治议会须与各行政区划的各个教会磋商：','（子）对自治议会的政策进行宣传，以提高其政策普及度：','（甲）提交定期报告；','（乙）维护完整记录；','第三条 进一步决定继续讨论该问题。'] if lang=='zh' else ['The European Commission,','Realizing the European Energy Security which served as an overall strategy to guard the EU’s energy safety,','Having devoted attention to current regional needs,','PART I Definition','1. Defines that an energy company shall comply with all existing regulations;','PART II Renewable Energy','2. Calls upon member states to cooperate:','(a) Encouraging renewable energy projects:','(i) Including solar generation;','(ii) Including wind generation;','i. Keeping a transparent reporting mechanism;','3. Further recommends the exchange of public information.']
            else: lines += ['1. 加入行动性条款于第五条第（一）款（子）：要求在英国结束委任统治后，即刻由委员会主持召开制宪会议；','2. 修改行动性条款原第五条第（一）款（丑）：自治委员会下设各省级市级和县级行政区的民族自治议会作为本地权力机构；','3. 删除行动性条款原第五条第（一）款（寅）。'] if lang=='zh' else ['1. Add as the operative clauses No. 6 entitled as Regional Cooperation including:','(a) Unification of energy policies including but not limited to:','(i) Policies concerning the promotion of usage of different forms of energy;','(ii) Policies concerning the price and number of imports and exports of energy resources;','2. Add as the operative clauses 9.b.V.: Allowing the setting up of renewable energy acceleration zone.']
        doc=Document(); section=doc.sections[0]; section.page_width=Inches(12); section.page_height=Inches(8); section.left_margin=Inches(4)
        doc.styles['Normal'].font.size=Pt(38)
        for number,text in enumerate(lines):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(3); p.paragraph_format.space_after=Pt(40)
            for i in range(0,len(text),2):
                run=p.add_run(text[i:i+2]); run.font.size=Pt(3 if i%4 else 35); run.bold=True; run.italic=True; run.underline=True
                fonts=run._r.get_or_add_rPr().get_or_add_rFonts(); fonts.set(qn('w:ascii'),'Comic Sans MS'); fonts.set(qn('w:eastAsia'),'Arial')
        name=f'{kind}-{lang}.docx'; doc.save(out/name); manifest.append({'file':name,'type':kind,'language':lang})
(out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(out)
