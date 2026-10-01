#!/usr/bin/env python3
"""Run formatting-only DOCX torture cases through a selected pipeline."""

from __future__ import annotations

import argparse
import json
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.docx_package import validate_docx_package  # noqa: E402
from app.fingerprint import canonical_body, visible_text_signature  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402


RUN_NOISE = (
    "vanish", "webHidden", "strike", "dstrike", "outline", "shadow", "emboss",
    "imprint", "caps", "smallCaps", "vertAlign", "position", "spacing", "w",
    "kern", "color", "highlight", "shd", "bdr", "effect", "glow", "reflection",
)
PARAGRAPH_NOISE = ("bidi", "textDirection", "shd", "pBdr", "framePr", "contextualSpacing", "snapToGrid")
SECTION_NOISE = ("cols", "lnNumType", "pgBorders", "docGrid")


def noise_count(document) -> int:
    count = 0
    for section in document.sections:
        count += sum(section._sectPr.find(qn(f"w:{tag}")) is not None for tag in SECTION_NOISE)
    for paragraph in document.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is not None:
            count += sum(ppr.find(qn(f"w:{tag}")) is not None for tag in PARAGRAPH_NOISE)
        for run in paragraph.runs:
            rpr = run._r.rPr
            if rpr is not None:
                count += sum(rpr.find(qn(f"w:{tag}")) is not None for tag in RUN_NOISE)
    return count


def page_summary(document) -> list[dict]:
    return [
        {
            "width_mm": round(section.page_width.mm, 2),
            "height_mm": round(section.page_height.mm, 2),
            "top_mm": round(section.top_margin.mm, 2),
            "bottom_mm": round(section.bottom_margin.mm, 2),
            "left_mm": round(section.left_margin.mm, 2),
            "right_mm": round(section.right_margin.mm, 2),
        }
        for section in document.sections
    ]


def run_case(document_type: str, correct: Path, stress: Path, output_dir: Path) -> dict:
    correct_bytes = correct.read_bytes()
    stress_bytes = stress.read_bytes()
    correct_doc = Document(BytesIO(correct_bytes))
    stress_doc = Document(BytesIO(stress_bytes))
    pipeline = PIPELINES[document_type](ROOT / "templates" / "pkunmun2026")
    result = pipeline.run(stress_bytes, normalize_punctuation=False, preserve_country_order=True)
    validate_docx_package(result.content)
    output = output_dir / f"{correct.stem}_程序修复.docx"
    output.write_bytes(result.content)
    repaired_doc = Document(BytesIO(result.content))
    errors = [item.to_dict() for item in result.validations if item.status == "error"]
    warnings = [item.to_dict() for item in result.validations if item.status == "warning"]
    text_preserved = visible_text_signature(correct_doc) == visible_text_signature(repaired_doc)
    canonical_preserved = canonical_body(correct_doc) == canonical_body(repaired_doc)
    remaining_noise = noise_count(repaired_doc)
    # The strict signature is stronger for this formatting-only test.  The
    # canonical body intentionally skips some header/continuation paragraphs
    # and can therefore change when a high-confidence structural repair only
    # changes paragraph boundaries.
    passed = not errors and text_preserved and remaining_noise == 0
    return {
        "document_type": document_type,
        "correct": str(correct),
        "stress": str(stress),
        "output": str(output),
        "passed": passed,
        "strict_visible_text_preserved": text_preserved,
        "canonical_body_preserved": canonical_preserved,
        "stress_noise_count": noise_count(stress_doc),
        "remaining_noise_count": remaining_noise,
        "page": page_summary(repaired_doc),
        "paragraphs": len(repaired_doc.paragraphs),
        "preambulatory_clauses": len(result.model.preambulatory_clauses),
        "operative_clauses": len(result.model.operative_clauses),
        "repair_actions": result.model.repair_actions,
        "validation_errors": errors,
        "validation_warnings": warnings,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--document-type", choices=sorted(PIPELINES), default="draft-resolution")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--case", action="append", nargs=2, metavar=("CORRECT", "STRESS"), required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    report = [
        run_case(args.document_type, Path(correct), Path(stress), args.output_dir)
        for correct, stress in args.case
    ]
    report_path = args.output_dir / "stress-regression.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(report_path)
    if not all(item["passed"] for item in report):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
