#!/usr/bin/env python3
"""Create deterministic formatting-only torture copies of correct resolutions.

The source files are opened read-only.  Paragraph text and paragraph order are
preserved exactly; only page, paragraph, style and run formatting are damaged.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT_NAMES = ("Comic Sans MS", "Courier New", "Arial Black", "Papyrus")
FONT_SIZES = (6, 9, 18, 28)
COLORS = (RGBColor(255, 0, 0), RGBColor(0, 90, 255), RGBColor(0, 150, 70), RGBColor(160, 0, 180))
HIGHLIGHTS = (WD_COLOR_INDEX.YELLOW, WD_COLOR_INDEX.BRIGHT_GREEN, WD_COLOR_INDEX.PINK, WD_COLOR_INDEX.TURQUOISE)
ALIGNMENTS = (
    WD_ALIGN_PARAGRAPH.CENTER,
    WD_ALIGN_PARAGRAPH.RIGHT,
    WD_ALIGN_PARAGRAPH.DISTRIBUTE,
    WD_ALIGN_PARAGRAPH.LEFT,
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def add_on_off(parent, tag: str) -> None:
    node = OxmlElement(f"w:{tag}")
    node.set(qn("w:val"), "1")
    parent.append(node)


def corrupt(source: Path, output: Path) -> dict:
    document = Document(source)
    original_text = [paragraph.text for paragraph in document.paragraphs]

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    if "Stress Chaos" not in document.styles:
        chaos = document.styles.add_style("Stress Chaos", WD_STYLE_TYPE.PARAGRAPH)
    else:
        chaos = document.styles["Stress Chaos"]
    chaos.font.name = "Impact"
    chaos.font.size = Pt(30)
    chaos.font.bold = True

    for section in document.sections:
        width, height = section.page_width, section.page_height
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = max(width, height)
        section.page_height = min(width, height)
        section.top_margin = Inches(3.8)
        section.bottom_margin = Inches(3.8)
        section.left_margin = Inches(4.2)
        section.right_margin = Inches(4.2)
        section.header_distance = Inches(2.2)
        section.footer_distance = Inches(2.2)
        sect_pr = section._sectPr
        cols = sect_pr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sect_pr.append(cols)
        cols.set(qn("w:num"), "3")
        cols.set(qn("w:space"), "720")
        cols.set(qn("w:sep"), "1")
        grid = OxmlElement("w:docGrid")
        grid.set(qn("w:type"), "linesAndChars")
        grid.set(qn("w:linePitch"), "80")
        sect_pr.append(grid)

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text
        paragraph.clear()
        paragraph.style = chaos if index % 3 == 0 else normal
        for char_index, char in enumerate(text):
            run = paragraph.add_run(char)
            variant = (index + char_index) % len(FONT_NAMES)
            run.font.name = FONT_NAMES[variant]
            run.font.size = Pt(FONT_SIZES[variant])
            run.font.color.rgb = COLORS[variant]
            run.font.highlight_color = HIGHLIGHTS[variant]
            run.bold = variant in (0, 3)
            run.italic = variant in (1, 3)
            run.underline = variant in (2, 3)
            rpr = run._r.get_or_add_rPr()
            if (index + char_index) % 11 == 0:
                add_on_off(rpr, "strike")
            if (index + char_index) % 13 == 0:
                add_on_off(rpr, "smallCaps")
            if (index + char_index) % 17 == 0:
                position = OxmlElement("w:position")
                position.set(qn("w:val"), "18")
                rpr.append(position)
            if (index + char_index) % 19 == 0:
                spacing = OxmlElement("w:spacing")
                spacing.set(qn("w:val"), "160")
                rpr.append(spacing)
            if (index + char_index) % 23 == 0:
                add_on_off(rpr, "vanish")

        paragraph.alignment = ALIGNMENTS[index % len(ALIGNMENTS)]
        paragraph.paragraph_format.left_indent = Inches(5.5 if index % 2 else -1.2)
        paragraph.paragraph_format.right_indent = Inches(3.0)
        paragraph.paragraph_format.first_line_indent = Inches(2.8 if index % 2 else -2.8)
        paragraph.paragraph_format.space_before = Pt(48 + index % 4 * 18)
        paragraph.paragraph_format.space_after = Pt(42 + index % 5 * 16)
        paragraph.paragraph_format.line_spacing = 4.0 if index % 2 else 0.6
        paragraph.paragraph_format.page_break_before = index > 0 and index % 3 == 0
        paragraph.paragraph_format.keep_together = index % 2 == 0
        paragraph.paragraph_format.keep_with_next = index % 2 == 1
        ppr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), ("FFCC00", "00CCFF", "FF99CC")[index % 3])
        ppr.append(shading)
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "double")
        bottom.set(qn("w:sz"), "24")
        bottom.set(qn("w:color"), "FF0000")
        borders.append(bottom)
        ppr.append(borders)
        if index % 4 == 0:
            add_on_off(ppr, "bidi")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    reopened = Document(output)
    output_text = [paragraph.text for paragraph in reopened.paragraphs]
    if output_text != original_text:
        raise RuntimeError(f"stress copy changed paragraph text: {source.name}")
    return {
        "source": str(source),
        "output": str(output),
        "source_sha256": sha256(source),
        "output_sha256": sha256(output),
        "paragraph_count": len(original_text),
        "text_preserved": output_text == original_text,
        "corruptions": [
            "landscape-and-invalid-margins",
            "three-columns-and-document-grid",
            "fragmented-character-runs",
            "mixed-fonts-sizes-colors-highlights",
            "strike-smallcaps-raised-hidden-expanded-text",
            "extreme-indents-spacing-and-forced-page-breaks",
            "mixed-alignment-bidi-shading-and-borders",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("sources", nargs="+", type=Path)
    args = parser.parse_args()
    manifest = []
    for source in args.sources:
        output = args.output_dir / f"{source.stem}_极限混乱.docx"
        manifest.append(corrupt(source, output))
    manifest_path = args.output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(manifest_path)


if __name__ == "__main__":
    main()
