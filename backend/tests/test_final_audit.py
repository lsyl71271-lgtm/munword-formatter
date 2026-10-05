"""Final-audit reproductions; no real user files or step-03 fixture repair."""
import sys
import asyncio
import struct
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
from app.docx_package import validate_docx_package
from app.errors import InvalidDocxError, InvalidRequestError, ProtectedContentError
from app.pipelines import PIPELINES
from app.content_guard import verify_package
from cli import write_atomic
from fastapi.testclient import TestClient
from app.main import app
from app.upload_limit import UploadLimitMiddleware


def source():
    doc = Document()
    for text in ("工作文件", "委员会：大会", "议题：合作", "国家：中国", "1. 支持合作。"):
        doc.add_paragraph(text)
    return save(doc)


def save(doc):
    stream = BytesIO(); doc.save(stream); return stream.getvalue()


def add_part(data, name, payload, compression=zipfile.ZIP_STORED):
    stream = BytesIO(data)
    with zipfile.ZipFile(stream, "a") as archive:
        archive.writestr(name, payload, compress_type=compression)
    return stream.getvalue()


class FinalAuditTests(unittest.TestCase):
    def test_corrupt_deflate_is_a_file_error_not_an_internal_error(self):
        data = bytearray(source())
        name_length, extra_length = struct.unpack_from("<HH", data, 26)
        start = 30 + name_length + extra_length
        data[start:start + 2] = b"\xff\xff"
        with self.assertRaises(InvalidDocxError):
            validate_docx_package(bytes(data))

    def test_local_zip_flags_must_match_the_central_directory(self):
        data = bytearray(source())
        struct.pack_into("<H", data, 6, struct.unpack_from("<H", data, 6)[0] | 1)
        with self.assertRaises(InvalidDocxError):
            validate_docx_package(bytes(data))

    def test_invalid_encoding_in_an_unused_xml_part_is_not_preserved(self):
        with self.assertRaises(InvalidDocxError):
            validate_docx_package(add_part(source(), "customXml/extra.xml", b"<x>\xff</x>"))

    def test_chunked_upload_is_bounded_before_the_handler_reads_it(self):
        handled, sent = [], []

        async def handler(scope, receive, send):
            handled.append(True)

        async def check():
            messages = iter([{"type": "http.request", "body": b"abc", "more_body": True},
                             {"type": "http.request", "body": b"def", "more_body": False}])

            async def receive():
                return next(messages)

            async def send(message):
                sent.append(message)

            await UploadLimitMiddleware(handler, max_bytes=5)(
                {"type": "http", "method": "POST", "path": "/api/parse/working-paper", "headers": []}, receive, send)

        asyncio.run(check())
        self.assertFalse(handled)
        self.assertEqual(sent[0]["status"], 413)

    def test_api_rejects_oversized_declared_body_before_multipart_parsing(self):
        with TestClient(app, base_url="http://localhost") as client:
            response = client.post("/api/parse/working-paper", files={"file": ("input.docx", source())},
                                   headers={"Content-Length": "999999999"})
        self.assertEqual(response.status_code, 413)

    def test_origin_guard_uses_asgi_path_not_a_host_poisoned_url(self):
        with TestClient(app, base_url="http://localhost") as client:
            response = client.post("/api/parse/working-paper", files={"file": ("input.docx", source())},
                                   headers={"Host": "localhost:8000/not-api", "Origin": "https://example.invalid"})
        self.assertEqual(response.status_code, 403)

    def test_atomic_cli_output_never_overwrites_an_existing_file_by_default(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "result.docx"
            target.write_bytes(b"keep this")
            with self.assertRaises(FileExistsError):
                write_atomic(target, b"new result")
            self.assertEqual(target.read_bytes(), b"keep this")
            self.assertEqual(list(Path(directory).iterdir()), [target])
            write_atomic(target, b"explicit replace", overwrite=True)
            self.assertEqual(target.read_bytes(), b"explicit replace")

    def test_duplicate_relationship_ids_cannot_hide_a_changed_target(self):
        namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
        rels = f'<Relationships xmlns="{namespace}"><Relationship Id="x" Target="old"/><Relationship Id="x" Target="unchanged"/></Relationships>'.encode()
        with self.assertRaises(InvalidDocxError):
            verify_package(add_part(source(), "word/_rels/extra.xml.rels", rels), source())

    def test_unsafe_paths_and_unsupported_compression_are_rejected(self):
        for name, method in (("word\\unsafe.bin", zipfile.ZIP_STORED), ("word/..", zipfile.ZIP_STORED),
                             ("word/extra.bin", zipfile.ZIP_BZIP2)):
            with self.subTest(name=name, method=method), self.assertRaises(InvalidDocxError):
                validate_docx_package(add_part(source(), name, b"data", method))

    def test_xml_limit_is_checked_before_inflation(self):
        with patch("app.docx_package.MAX_XML_BYTES", 100):
            with self.assertRaisesRegex(InvalidDocxError, "XML"):
                validate_docx_package(source())

    def test_direct_pipeline_and_cli_cannot_bypass_package_safety(self):
        data = add_part(source(), "word/extra.xml", b'<!DOCTYPE x [<!ENTITY a "b">]><x/>')
        with self.assertRaises(InvalidDocxError):
            PIPELINES["working-paper"](ROOT / "templates/pkunmun2026").run(data)

    def test_step03_rejects_invalid_xml_characters(self):
        pipeline = PIPELINES["working-paper"](ROOT / "templates/pkunmun2026")
        for value in ("A\0B", "A\uffffB", "A\ud800B"):
            with self.subTest(value=repr(value)), self.assertRaises(InvalidRequestError):
                pipeline.run(source(), overrides={"topic": value})

    def test_manual_complex_scalar_field_never_silently_ignores_the_requested_edit(self):
        doc = Document(BytesIO(source()))
        run = doc.paragraphs[3].runs[0]
        run._r.append(OxmlElement("w:drawing"))
        with self.assertRaises(ProtectedContentError):
            PIPELINES["working-paper"](ROOT / "templates/pkunmun2026").run(save(doc), overrides={"country": "法国"})

    def test_position_body_label_does_not_move_the_header_boundary(self):
        doc = Document()
        for text in ("立场文件", "委员会：大会", "议题：合作", "国家：中国", "代表：测试", "（一）背景",
                     "1. 支持各国合作。", "议题：正文引用的标签，不是页首。"):
            doc.add_paragraph(text)
        result = PIPELINES["position-paper"](ROOT / "templates/pkunmun2026").run(save(doc))
        texts = [p.text for p in Document(BytesIO(result.content)).paragraphs]
        self.assertIn("1、\t支持各国合作。", texts)
        self.assertIn("议题：正文引用的标签，不是页首。", texts)


if __name__ == "__main__":
    unittest.main()
