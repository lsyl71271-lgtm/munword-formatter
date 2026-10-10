import sys
import unittest
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
from app.pipelines import PIPELINES
from app.treaty import is_treaty
from app.numbering import sequence_issues
from app.content_guard import verify_package

# The handbook's six types; diplomatic agreements and joint statements have their own tests.
HANDBOOK_TYPES = [kind for kind in PIPELINES if not is_treaty(kind)]

TITLES = dict(zip(PIPELINES,['立场文件','工作文件','指令草案','决议草案','友好修正案','非友好修正案']))

def source(kind,lines,en=False):
    d=Document()
    titles=['Position Paper','Working Paper','Draft Directive','Draft Resolution','Friendly Amendment','Unfriendly Amendment']
    header=['Committee: General Assembly','Topic: Cooperation','Country: China','Sponsors: China','Signatories: France'] if en else ['委员会：联合国大会','议题：合作','国家：中国','起草国：中国','附议国：法国']
    for t in [titles[HANDBOOK_TYPES.index(kind)] if en else TITLES[kind],*header,*lines]:
        d.add_paragraph(t)
    b=BytesIO();d.save(b);return b.getvalue()

def run(kind,data):
    r=PIPELINES[kind](ROOT/'templates/pkunmun2026').run(data)
    assert not [v for v in r.validations if v.status=='error'],r.validations
    return r,Document(BytesIO(r.content))

class AllNumberingTests(unittest.TestCase):
    def test_all_types_check_even_without_numbering(self):
        for kind in HANDBOOK_TYPES:
            with self.subTest(kind=kind):
                r,d=run(kind,source(kind,['普通正文提到第一条和（一），不应改写。']))
                self.assertTrue(any(v.code=='numbering-policy' and '未发现' in v.detail for v in r.validations))
                self.assertIn('普通正文提到第一条和（一），不应改写。',[p.text for p in d.paragraphs])

    def test_working_parts_and_references(self):
        r,d=run('working-paper',source('working-paper',['Part I: Cooperation','第三条 提交方案：','（一）保留“第一条”引用。','Part II: Review','第一条 审查。']))
        texts=[p.text for p in d.paragraphs]
        for text in ['3. 提交方案：','(a)保留“第一条”引用。','1. 审查。']: self.assertIn(text,texts)
        r2,d2=run('working-paper',r.content)
        self.assertEqual(texts,[p.text for p in d2.paragraphs])

    def test_directive_profiles(self):
        for en in (False,True):
            lines=['1. Requests cooperation:','(a) Provide resources.'] if en else ['第一条 要求合作：','(a) 提交方案。']
            r,d=run('draft-directive',source('draft-directive',lines,en))
            texts=[p.text for p in d.paragraphs]
            self.assertTrue(any(t.startswith('a. Provide' if en else '（一） 提交') for t in texts))
            self.assertTrue(any(t.startswith('1. Requests' if en else '1. 要求') for t in texts))
            r2,d2=run('draft-directive',r.content)
            self.assertEqual(texts,[p.text for p in d2.paragraphs])

    def test_position_sections_and_proposals(self):
        r,d=run('position-paper',source('position-paper',['（一）议题背景','普通论述。','一、国际措施','1. 建议：','(a) 提供资源。']))
        texts=[p.text for p in d.paragraphs]
        for text in ['（一）议题背景','一、国际措施','a) 提供资源。']: self.assertIn(text,texts)
        r2,d2=run('position-paper',r.content)
        self.assertEqual(texts,[p.text for p in d2.paragraphs])

    def test_amendment_only_operation_prefix(self):
        for kind in ('friendly-amendment','unfriendly-amendment'):
            r,d=run(kind,source(kind,['第一条 修改第五条第（一）款（子）：“原文”；','2. 加入如下条款：','（a）引用的待新增内容。']))
            texts=[p.text for p in d.paragraphs]
            self.assertTrue(any(t.startswith('1. 修改第五条第（一）款（子）：“原文”') for t in texts))
            self.assertTrue(any(t.startswith('（a）引用的待新增内容') for t in texts))
            self.assertTrue(any(v.code=='list-numbering' and '目标草案' in v.detail for v in r.validations))
            r2,d2=run(kind,r.content)
            self.assertEqual(texts,[p.text for p in d2.paragraphs])

    def test_non_dr_native_counter_protection(self):
        d=Document(BytesIO(source('working-paper',['1. 方案：'])))
        p=d.add_paragraph('资源。');npr=p._p.get_or_add_pPr().get_or_add_numPr();npr.get_or_add_numId().val=1;npr.get_or_add_ilvl().val=0
        n=d.part.numbering_part.element
        num=next(t for t in n.findall(qn('w:num')) if t.get(qn('w:numId'))=='1')
        aid=num.find(qn('w:abstractNumId')).get(qn('w:val'))
        lvl=next(t for t in n.findall(qn('w:abstractNum')) if t.get(qn('w:abstractNumId'))==aid).find(qn('w:lvl'))
        for tag,value in [('numFmt','chineseCounting'),('lvlText','（%1）'),('start','3')]: lvl.find(qn('w:'+tag)).set(qn('w:val'),value)
        b=BytesIO();d.save(b);before=b.getvalue()
        r,d2=run('working-paper',before)
        new=next(t for t in d2.part.numbering_part.element.findall(qn('w:num')) if t.get(qn('w:numId'))=='1').find(qn('w:lvlOverride')).find(qn('w:lvl'))
        self.assertEqual(new.find(qn('w:numFmt')).get(qn('w:val')),'lowerLetter')
        self.assertEqual(new.find(qn('w:start')).get(qn('w:val')),'3')
        rules=[dict(id='1',ilvl='0',level=1,language='zh',type='working-paper')]
        self.assertEqual(verify_package(before,r.content,rules),[])
        new.find(qn('w:start')).set(qn('w:val'),'4');b=BytesIO();d2.save(b)
        self.assertTrue(verify_package(before,b.getvalue(),rules))

    def test_ambiguity_is_not_guessed(self):
        r,d=run('working-paper',source('working-paper',['(a) 没有父项。']))
        self.assertIn('(a) 没有父项。',[p.text for p in d.paragraphs])
        self.assertTrue(any('无明确顶层' in v.detail for v in r.validations))

    def test_sequences_roman_extended_letters_and_restarts(self):
        items=[dict(text='1. Parent:',family='decimal',value=1,level=0),dict(text='(i) a',family='roman',value=1,level=2),dict(text='(iii) b',family='roman',value=3,level=2),dict(text='Part II',family='',value=None,level=0,reset=True),dict(text='1. Start again',family='decimal',value=1,level=0)]
        self.assertEqual(sequence_issues(items),[dict(index=2,previous=1)])

