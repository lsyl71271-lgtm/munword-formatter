"""Regression tests for non-text content that formatting used to delete.

Every case here passed every validation in v1.5.0 while silently dropping
content: the fingerprints only compared plain paragraph text, so an image, a
section break or a table was invisible to both the formatter and its checks.
"""

import base64
import sys
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app.docx_view import structure_losses, structure_signature  # noqa: E402
from app.ooxml_edit import edit_visible_text  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402

TEMPLATES = ROOT / "templates" / "pkunmun2026"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _png_path(tmp: Path) -> str:
    target = tmp / "pixel.png"
    target.write_bytes(PNG)
    return str(target)


def _header(document, title="指令草案"):
    for text in (title, "委员会：联合国大会", "起草国：日本国", "附议国：法兰西共和国", "联合国大会，"):
        document.add_paragraph(text)


def _save(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _run(document_type, content, **kwargs):
    return PIPELINES[document_type](TEMPLATES).run(content, **kwargs)


def _drawings(content: bytes) -> int:
    return Document(BytesIO(content)).element.body.xml.count("<w:drawing")


class ContentProtectionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(__file__).resolve().parent / "_tmp"
        self.tmp.mkdir(exist_ok=True)
        self.png = _png_path(self.tmp)

    def tearDown(self):
        for item in self.tmp.glob("*"):
            item.unlink()
        self.tmp.rmdir()

    def test_image_only_paragraph_survives_empty_paragraph_collapsing(self):
        document = Document()
        _header(document)
        document.add_paragraph("1. 决定设立观察团")
        document.add_paragraph().add_run().add_picture(self.png, width=Inches(1))
        result = _run("draft-directive", _save(document))

        self.assertEqual(_drawings(result.content), 1)
        self.assertNotIn("error", [item.status for item in result.validations])

    def test_inline_image_survives_punctuation_normalization(self):
        document = Document()
        _header(document)
        paragraph = document.add_paragraph("1. 决定设立观察团，并参考 ")
        paragraph.add_run("附图：")
        paragraph.add_run().add_picture(self.png, width=Inches(1))
        result = _run("draft-directive", _save(document), normalize_punctuation=True)

        self.assertEqual(_drawings(result.content), 1)
        texts = [item.text for item in Document(BytesIO(result.content)).paragraphs]
        self.assertTrue(any(text.endswith("。") or text.endswith("；") for text in texts))
        self.assertNotIn("error", [item.status for item in result.validations])

    def test_source_emphasis_inside_a_clause_is_not_flattened(self):
        # Rewriting a proposal's marker ("1." -> "1、\t") must edit in place.
        # Position papers keep the author's emphasis (handbook 页17–18); the
        # other document types clear clause emphasis by design.
        document = Document()
        for text in ("立场文件", "委员会：联合国大会", "议题：海洋治理", "国家：日本国", "代表：示例"):
            document.add_paragraph(text)
        paragraph = document.add_paragraph("1. 要求秘书长")
        emphasized = paragraph.add_run("按季度")
        emphasized.bold = True
        paragraph.add_run("提交报告")
        result = _run("position-paper", _save(document), normalize_punctuation=True)

        output = Document(BytesIO(result.content))
        bold_runs = [run.text for item in output.paragraphs for run in item.runs if run.bold]
        self.assertIn("按季度", bold_runs)

    def test_section_break_is_preserved(self):
        document = Document()
        _header(document)
        document.add_paragraph("1. 决定设立观察团")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("2. 要求秘书长报告")
        result = _run("draft-directive", _save(document))

        self.assertEqual(len(Document(BytesIO(result.content)).sections), 2)
        self.assertNotIn("error", [item.status for item in result.validations])

    def test_table_text_is_formatted_and_fingerprinted(self):
        document = Document()
        _header(document, title="工作文件")
        document.add_paragraph("1. 建议建立信息共享机制")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "表格中的正文内容"
        table.cell(0, 1).text = "补充说明"
        result = _run("working-paper", _save(document))

        output = Document(BytesIO(result.content))
        self.assertEqual(len(output.tables), 1)
        sizes = {
            run.font.size.pt
            for cell in output.tables[0].rows[0].cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if run.text.strip() and run.font.size is not None
        }
        self.assertTrue(sizes, "table runs must carry an explicit size")
        self.assertEqual(len(sizes), 1)

    def test_friendly_amendment_keeps_image_in_numbered_operation(self):
        document = Document()
        _header(document, title="友好修正案")
        paragraph = document.add_paragraph("1. 修改第一条：增加附图 ")
        paragraph.add_run().add_picture(self.png, width=Inches(1))
        result = _run("friendly-amendment", _save(document))

        self.assertEqual(_drawings(result.content), 1)
        self.assertNotIn("error", [item.status for item in result.validations])
        output = Document(BytesIO(result.content))
        operation = next(item for item in output.paragraphs if "修改第一条" in item.text)
        # The typed operation number stays as written (no conversion to Word numbering).
        self.assertEqual(operation.text.strip(), "1. 修改第一条：增加附图")
        self.assertIsNone(operation._p.pPr.numPr)
        # Only the operation verb is italic (handbook 页52–53).
        self.assertEqual([run.text.strip() for run in operation.runs if run.italic], ["修改"])

    def test_country_continuation_survives_a_refused_list_rewrite(self):
        # The first list line holds a tracked insertion, so the list cannot be
        # rewritten inline; its plain continuation line used to be cleared
        # anyway, deleting a country.
        document = Document()
        for text in ("决议草案1.0", "委员会：联合国大会", "起草国：日本国"):
            document.add_paragraph(text)
        label = document.add_paragraph("附议国：")
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "1")
        insertion.set(qn("w:author"), "A")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "加拿大、澳大利亚、"
        run.append(text)
        insertion.append(run)
        label._p.append(insertion)
        document.add_paragraph("乍得共和国")
        for text in ("联合国大会，", "回顾以往决议，", "第一条 决定设立观察团。"):
            document.add_paragraph(text)
        result = _run("draft-resolution", _save(document), normalize_punctuation=True)

        self.assertNotIn("error", [item.status for item in result.validations])
        output = "".join(item.text for item in Document(BytesIO(result.content)).paragraphs)
        self.assertIn("乍得共和国", output)

    def test_tracked_insertion_is_never_dropped_silently(self):
        # A tracked insertion is not part of ``paragraph.text``; rebuilding the
        # paragraph used to delete it while every validation reported pass.
        document = Document()
        _header(document, title="友好修正案")
        paragraph = document.add_paragraph("1. 修改第一条：")
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "1")
        insertion.set(qn("w:author"), "A")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "新增的修订文字"
        run.append(text)
        insertion.append(run)
        paragraph._p.append(insertion)
        result = _run("friendly-amendment", _save(document))

        self.assertIn("新增的修订文字", Document(BytesIO(result.content)).element.body.xml)
        self.assertNotIn("error", [item.status for item in result.validations])

    def test_structure_signature_counts_tracked_changes(self):
        document = Document()
        paragraph = document.add_paragraph("正文")
        paragraph._p.append(OxmlElement("w:ins"))
        signature = structure_signature(document)
        paragraph._p.remove(paragraph._p[-1])

        self.assertEqual(signature["w:ins"], 1)
        self.assertEqual(structure_losses(signature, structure_signature(document)), ["修订（插入） 减少 1 处"])

    def test_structure_signature_counts_non_text_features(self):
        document = Document()
        document.add_paragraph("正文")
        document.add_paragraph().add_run().add_picture(self.png, width=Inches(1))
        document.add_table(rows=1, cols=1)
        signature = structure_signature(document)

        self.assertEqual(signature["w:drawing"], 1)
        self.assertEqual(signature["w:tbl"], 1)
        self.assertGreaterEqual(signature["w:sectPr"], 1)

    def test_edit_visible_text_keeps_other_nodes(self):
        document = Document()
        paragraph = document.add_paragraph("1. 决定设立观察团")
        paragraph.add_run().add_picture(self.png, width=Inches(1))
        before = document.element.body.xml.count("<w:drawing")

        self.assertTrue(edit_visible_text(paragraph, "决定设立观察团；"))
        self.assertEqual(paragraph.text, "决定设立观察团；")
        self.assertEqual(document.element.body.xml.count("<w:drawing"), before)

    def test_edit_visible_text_refuses_to_cross_a_line_break(self):
        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run("第一行")
        run.add_break()
        run.add_text("第二行")

        self.assertFalse(edit_visible_text(paragraph, "完全不同的内容"))
        self.assertEqual(paragraph.text, "第一行\n第二行")


if __name__ == "__main__":
    unittest.main()
