import sys
import unittest
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'backend'))
from app.pipelines import PIPELINES
from app.numbering import valid_marker_change, plan_hierarchy, marker_of, marker_text, native_range_reason
from app.content_guard import verify_package

def source(lines):
    d=Document()
    for t in ['决议草案1.1','联合国大会','议题：测试','起草国：中国','回顾原则，',*lines]: d.add_paragraph(t)
    b=BytesIO();d.save(b);return b.getvalue()
def run(data):
    r=PIPELINES['draft-resolution'](ROOT/'templates/pkunmun2026').run(data)
    assert not [v for v in r.validations if v.status=='error'],r.validations
    return r,Document(BytesIO(r.content))

class DrNumberingTests(unittest.TestCase):
    def test_ambiguity_survives_second_pass(self):
        r,d=run(source(['第一条 方案：','（一）机构：','（一）秘书处：','（二）监督。']))
        r2,d2=run(r.content)
        self.assertIn('（一）机构：',[p.text for p in d.paragraphs])
        for result in (r,r2):
            self.assertTrue(any(v.code=='dr-numbering' and '冲突' in v.detail for v in result.validations))
        self.assertEqual([p.text for p in d.paragraphs],[p.text for p in d2.paragraphs])

    def test_canonical_siblings_after_grandchild(self):
        expected=['第一条 方案：','（一）机构：','（子）秘书处：','（甲）存档；','（丑）审查；','（二）监督。']
        r,d=run(source(expected))
        self.assertEqual([p.text for p in d.paragraphs if p.text.startswith(('第','（'))],expected)
        self.assertFalse(any(v.code=='dr-numbering' and '冲突' in v.detail for v in r.validations))

    def test_jumps_not_renumbered(self):
        r,d=run(source(['第三条 方案：','（c）记录；','（e）核查。']))
        self.assertEqual([p.text for p in d.paragraphs if p.text.startswith(('第','（'))],['第三条 方案：','（三）记录；','（五）核查。'])

    def test_native_identities_and_missing_parent_evidence(self):
        items=[dict(text='第一条 方案：',family='article',value=1,level=0,top=True),dict(text='机构：',family='letter',value=None,level=1,key='1:0'),dict(text='秘书处。',family='letter',value=None,level=1,key='2:0'),dict(text='监督。',family='letter',value=None,level=1,key='1:0')]
        self.assertEqual([p['level'] for p in plan_hierarchy(items,'zh')],[0,1,2,1])
        items=[dict(text='第一条 方案。',family='article',value=1,level=0,top=True),dict(text='(a)机构；',family='letter',value=1,level=1),dict(text='（一）秘书处；',family='chinese',value=1,level=1)]
        self.assertIn('缺少明确父子关系',plan_hierarchy(items,'zh')[2]['reason'])

    def test_extended_and_english_markers(self):
        self.assertEqual([marker_text('en',level,4) for level in range(4)],['4.','(d)','(iv)','iv.'])
        self.assertEqual(marker_text('zh',2,13),'（子子）')
        self.assertEqual(marker_text('zh',3,11),'（甲甲）')
        self.assertEqual(marker_of('（子子）内容')['value'],13)
        self.assertEqual(marker_of('(i)内容')['value'],9)
        self.assertEqual(marker_of('(i)内容',True)['value'],1)

    def test_native_cycle_bounds(self):
        from lxml import etree
        xml=etree.fromstring('<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="12"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="(%1)"/></w:lvl></w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>')
        rule=dict(id='1',ilvl='0',level=2,language='zh')
        self.assertEqual(native_range_reason(xml,rule,1),'')
        self.assertIn('循环格式',native_range_reason(xml,rule,2))

    def test_missing_subject_not_preamble(self):
        r,d=run(source(['第一条呼吁合作：','（a）设立机构。']))
        self.assertIn('第一条呼吁合作：',[p.text for p in d.paragraphs])
        self.assertIn('（一）设立机构。',[p.text for p in d.paragraphs])

    def test_foreign_context_and_idempotence(self):
        r,d=run(source(['第一条 要求合作：','（a）建立机构：','（一）建立秘书处：','（甲）保存原始记录；','（二）提交报告；','（b）监督执行。']))
        body=[p.text for p in d.paragraphs if p.text.startswith(('第','（'))]
        self.assertEqual(body,['第一条 要求合作：','（一）建立机构：','（子）建立秘书处：','（甲）保存原始记录；','（丑）提交报告；','（二）监督执行。'])
        self.assertEqual([round(p.paragraph_format.left_indent.pt) for p in d.paragraphs if p.text.startswith(('第','（'))],[0,48,68,87,68,48])
        again,d2=run(r.content)
        self.assertEqual([p.text for p in d.paragraphs],[p.text for p in d2.paragraphs])
        self.assertTrue(any(v.code=='dr-numbering' for v in r.validations))

    def test_native_start_restart_and_guard(self):
        data=source(['第一条 要求合作：'])
        d=Document(BytesIO(data));p=d.add_paragraph('提交报告。')
        numpr=p._p.get_or_add_pPr().get_or_add_numPr();numpr.get_or_add_numId().val=1;numpr.get_or_add_ilvl().val=0
        n=d.part.numbering_part.element
        num=next(t for t in n.findall(qn('w:num')) if t.get(qn('w:numId'))=='1')
        abstract_id=num.find(qn('w:abstractNumId')).get(qn('w:val'))
        a=next(t for t in n.findall(qn('w:abstractNum')) if t.get(qn('w:abstractNumId'))==abstract_id)
        lvl=a.find(qn('w:lvl'));lvl.find(qn('w:numFmt')).set(qn('w:val'),'lowerLetter');lvl.find(qn('w:lvlText')).set(qn('w:val'),'(%1)');lvl.find(qn('w:start')).set(qn('w:val'),'3')
        b=BytesIO();d.save(b);before=b.getvalue()
        r,d2=run(before)
        n2=d2.part.numbering_part.element
        override=next(t for t in n2.findall(qn('w:num')) if t.get(qn('w:numId'))=='1').find(qn('w:lvlOverride'))
        self.assertEqual(override.find(qn('w:lvl')).find(qn('w:numFmt')).get(qn('w:val')),'chineseCounting')
        self.assertEqual(override.find(qn('w:lvl')).find(qn('w:start')).get(qn('w:val')),'3')
        rules=[dict(id='1',ilvl='0',level=1,language='zh')]
        self.assertEqual(verify_package(before,r.content,rules),[])
        override.find(qn('w:lvl')).find(qn('w:start')).set(qn('w:val'),'4')
        b=BytesIO();d2.save(b)
        self.assertTrue(verify_package(before,b.getvalue(),rules),'allowed display changes must not exempt counter tampering')

    def test_guard_and_ambiguity(self):
        self.assertTrue(valid_marker_change('(b)正文；','（二）正文。'))
        self.assertFalse(valid_marker_change('(b)正文；','（三）正文；'))
        self.assertFalse(valid_marker_change('(b)正文；','（二）更改正文；'))
        self.assertFalse(valid_marker_change('(i)正文','（一）正文'))
        self.assertTrue(valid_marker_change('(i)正文','（九）正文'))
        items=[dict(text='第一条 方案：',family='article',value=1,level=0,top=True),dict(text='（一）方案：',family='chinese',value=1,level=1),dict(text='（一）子项。',family='chinese',value=1,level=1)]
        self.assertIn('冲突',plan_hierarchy(items,'zh')[2]['reason'])
