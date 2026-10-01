"""Regressions from the independent v1.7.0 review (Claude).

Each case failed before its fix.  Expectations come from the OOXML structure
and the handbook rules, read back with plain lxml here, not from the engine.
"""

import sys
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app import content_guard  # noqa: E402
from app.docx_package import validate_docx_package  # noqa: E402
from app.errors import InvalidDocxError  # noqa: E402
from app.main import app  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402

TEMPLATES = ROOT / "templates" / "pkunmun2026"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = f'xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml"'
HEADER = ["决议草案1.0", "联合国大会", "测试议题", "起草国：法兰西共和国、日本国", "附议国：大韩民国", "联合国大会，", "回顾以往决议，"]


def q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def fragment(xml: str):
    return etree.fromstring(f"<w:wrap {NS}>{xml}</w:wrap>")[0]


def source(lines=HEADER) -> Document:
    document = Document()
    for text in lines:
        document.add_paragraph(text)
    return document


def save(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def run(document_type: str, content: bytes):
    return PIPELINES[document_type](TEMPLATES).run(content)


def errors(result):
    return [f"{item.label}: {item.detail}" for item in result.validations if item.status == "error"]


def body_xml(content: bytes):
    return etree.fromstring(zipfile.ZipFile(BytesIO(content)).read("word/document.xml"))


def own_text(paragraph) -> str:
    """Text of a paragraph without the paragraphs nested in its text boxes."""

    out = []
    for node in paragraph.iter(q("t")):
        ancestor = node.getparent()
        nested = False
        while ancestor is not paragraph:
            if ancestor.tag in (q("p"), q("del")):
                nested = True
                break
            ancestor = ancestor.getparent()
        if not nested:
            out.append(node.text or "")
    return "".join(out)


def italic_chars(paragraph) -> str:
    out = []
    for r in paragraph.iter(q("r")):
        if r.getparent() is not paragraph and r.getparent().tag != q("hyperlink"):
            continue
        rpr = r.find(q("rPr"))
        node = rpr.find(q("i")) if rpr is not None else None
        on = node is not None and node.get(q("val")) not in ("0", "false")
        if on:
            out.extend(t.text or "" for t in r.iter(q("t")))
    return "".join(out)


class TextBoxTests(unittest.TestCase):
    def test_ending_of_a_paragraph_holding_a_text_box_is_not_a_content_change(self):
        document = source()
        paragraph = document.add_paragraph("第一条 要求各国合作，见右框")
        paragraph._p.append(fragment(
            '<w:r><w:pict><v:shape style="width:120pt;height:40pt"><v:textbox><w:txbxContent>'
            '<w:p><w:r><w:t>（一）文本框内容，</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
        document.add_paragraph("第二条 决定继续审议。")
        result = run("draft-resolution", save(document))
        self.assertEqual(errors(result), [])
        body = body_xml(result.content)
        boxed = [p for p in body.iter(q("p")) if p.getparent().tag == q("txbxContent")]
        self.assertEqual([own_text(p) for p in boxed], ["（一）文本框内容，"], "text-box text is the author's")
        article = next(p for p in body.iter(q("p")) if own_text(p).startswith("第一条"))
        self.assertEqual(own_text(article), "第一条 要求各国合作，见右框；")


class TableTests(unittest.TestCase):
    def test_text_in_a_table_cell_is_not_turned_into_a_clause(self):
        document = source()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).paragraphs[0].add_run("第一条 表格中的条款；")
        document.add_paragraph("第二条 决定继续审议。")
        result = run("draft-resolution", save(document))
        self.assertEqual(errors(result), [])
        body = body_xml(result.content)
        cells = list(body.iter(q("tc")))
        self.assertEqual([[own_text(p) for p in cell.iter(q("p"))] for cell in cells], [["第一条 表格中的条款；"], [""]])
        for cell in cells:
            for paragraph in cell.iter(q("p")):
                self.assertIsNone(paragraph.find(f"{q('pPr')}/{q('numPr')}"))


class ContentControlTests(unittest.TestCase):
    def _source(self) -> bytes:
        document = source()
        document.element.body.insert(len(document.element.body) - 1, fragment(
            '<w:sdt><w:sdtPr><w:alias w:val="正文"/></w:sdtPr><w:sdtContent>'
            '<w:p><w:r><w:t xml:space="preserve">第一条 要求各国合作，</w:t></w:r></w:p>'
            '<w:p><w:r><w:t xml:space="preserve">第二条 决定继续审议</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        return save(document)

    def test_articles_in_a_content_control_follow_the_handbook(self):
        result = run("draft-resolution", self._source())
        self.assertEqual(errors(result), [])
        body = body_xml(result.content)
        inside = [p for p in body.iter(q("p")) if p.getparent().tag == q("sdtContent") and own_text(p)]
        self.assertEqual([own_text(p) for p in inside], ["第一条 要求各国合作；", "第二条 决定继续审议。"])
        self.assertEqual([italic_chars(p) for p in inside], ["要求", "决定"])
        self.assertEqual(len(list(body.iter(q("sdt")))), 1, "the control itself stays")

    def test_second_pass_adds_nothing_inside_a_content_control(self):
        first = run("draft-resolution", self._source())
        second = run("draft-resolution", first.content)
        texts = lambda content: [own_text(p) for p in body_xml(content).iter(q("p"))]
        self.assertEqual(texts(second.content), texts(first.content))

    def test_guard_sees_an_unlogged_change_inside_a_content_control(self):
        document = Document(BytesIO(self._source()))
        before = content_guard.Snapshot.take(document)
        text = next(node for node in document.element.body.iter(q("t")) if (node.text or "").startswith("第二条"))
        text.text = "第二条 决定停止审议"
        problems = content_guard.verify_format(before, document, {}, allowed_titles=(), labels=())
        self.assertTrue(problems)


class HeaderBoundaryTests(unittest.TestCase):
    """A body paragraph that starts like a header label is the author's text."""

    LINES = ["决议草案1.0", "联合国大会", "测试议题", "起草国：法兰西共和国", "联合国大会，", "回顾以往决议，",
             "第一条 决定设立工作组；", "议题：后续安排如下，", "第二条 决定继续审议。"]

    def test_label_in_the_body_is_neither_metadata_nor_dropped(self):
        content = save(source(self.LINES))
        model = PIPELINES["draft-resolution"](TEMPLATES).parse(content)
        self.assertEqual(model.topic, "测试议题")
        result = run("draft-resolution", content)
        self.assertEqual(errors(result), [])
        texts = [own_text(p) for p in body_xml(result.content).iter(q("p")) if own_text(p)]
        # The label stays; the clause ending follows the punctuation rule.
        self.assertTrue(any(text.startswith("议题：后续安排如下") for text in texts), texts)
        first = next(p for p in body_xml(result.content).iter(q("p")) if own_text(p).startswith("第一条"))
        self.assertEqual(italic_chars(first), "决定", "the first article is a clause, not a header line")

    def test_step_three_topic_never_rewrites_a_body_paragraph(self):
        content = save(source(self.LINES))
        result = PIPELINES["draft-resolution"](TEMPLATES).run(content, overrides={"topic": "新议题"})
        self.assertEqual(errors(result), [])
        texts = [own_text(p) for p in body_xml(result.content).iter(q("p")) if own_text(p)]
        self.assertIn("新议题", texts)
        self.assertTrue(any(text.startswith("议题：后续安排如下") for text in texts), texts)


class CountryListTests(unittest.TestCase):
    def test_english_subject_line_after_signatories_is_not_a_country(self):
        lines = ["Draft Directive 1.1", "Committee: General Assembly", "Sponsors: Japan, France", "Signatories: Canada, Republic of Korea",
                 "The Executive Council,", "1. Requests all Member States to submit annual reports;", "2. Decides to remain seized of the matter."]
        content = save(source(lines))
        model = PIPELINES["draft-directive"](TEMPLATES).parse(content)
        self.assertEqual(model.signatories, ["Canada", "Republic of Korea"])
        result = run("draft-directive", content)
        self.assertEqual(errors(result), [])
        texts = [own_text(p) for p in body_xml(result.content).iter(q("p")) if own_text(p)]
        self.assertIn("Signatories: Canada, Republic of Korea", texts)
        self.assertIn("The Executive Council,", texts)


def indent_of(paragraph) -> dict:
    ind = paragraph.find(f"{q('pPr')}/{q('ind')}")
    return {} if ind is None else {key.split("}")[-1]: value for key, value in ind.attrib.items()}


class SubjectLineTests(unittest.TestCase):
    def test_subject_line_named_with_the_is_recognized(self):
        lines = ["Draft Resolution 2.1", "Committee: Security Council", "Topic: Maritime security", "Sponsors: France, Japan",
                 "Signatories: Canada", "The Security Council,", "Recalling its previous resolutions,",
                 "1. Requests all States to cooperate;", "2. Decides to remain seized of the matter."]
        result = run("draft-resolution", save(source(lines)))
        self.assertEqual(errors(result), [])
        subject = next(p for p in body_xml(result.content).iter(q("p")) if own_text(p).startswith("The Security"))
        self.assertEqual(own_text(subject), "The Security Council,", "a subject keeps its comma")
        self.assertEqual(italic_chars(subject), "The Security Council,", "handbook: the subject line is italic")


class LevelTests(unittest.TestCase):
    """Levels from markers and numbering; indentation only inside a list."""

    def _damaged_list(self) -> bytes:
        document = source(["工作文件2.1", "联合国教科文组织", "文物保护", "起草国：法兰西共和国"])
        document.add_paragraph("1.建立文物登记制度；")
        document.add_paragraph("2.明确预防和打击文物走私的处理办法：")
        first = document.add_paragraph("建立国际文物数据库，加强海关识别能力；")  # lost its number, kept its indent
        first._p.get_or_add_pPr().append(fragment('<w:ind w:left="420"/>'))
        document.add_paragraph("3.鼓励各国交流。")
        return save(document)

    def test_a_damaged_item_is_nested_and_a_second_pass_changes_nothing(self):
        first = run("working-paper", self._damaged_list())
        self.assertEqual(errors(first), [])
        second = run("working-paper", first.content)
        read = lambda content: [(own_text(p), indent_of(p)) for p in body_xml(content).iter(q("p"))]
        self.assertEqual(read(second.content), read(first.content))
        item = next(p for p in body_xml(first.content).iter(q("p")) if own_text(p).startswith("建立国际"))
        self.assertEqual(indent_of(item).get("left"), "840", "nested one level under 2.")

    def test_a_stray_indent_outside_a_list_is_not_a_level(self):
        document = source(["工作文件2.1", "联合国教科文组织", "文物保护", "起草国：法兰西共和国"])
        stray = document.add_paragraph("各代表团就文物保护展开讨论，形成以下建议")
        stray._p.get_or_add_pPr().append(fragment('<w:ind w:left="840" w:hanging="420"/>'))
        document.add_paragraph("1.建立文物登记制度；")
        document.add_paragraph("2.鼓励各国交流。")
        result = run("working-paper", save(document))
        self.assertEqual(errors(result), [])
        paragraph = next(p for p in body_xml(result.content).iter(q("p")) if own_text(p).startswith("各代表团"))
        self.assertNotIn("hanging", indent_of(paragraph))


class MarkerTests(unittest.TestCase):
    LINES = ["工作文件3.1", "联合国教科文组织", "文物保护", "起草国：法兰西共和国",
             "5．建立地区互助机制，明确其运行方式：", "（a）决策一国一票；", "（b）资金透明；", "（c）尊重主权；", "（d）定期报告；",
             "（i）按年度提交；", "（ii）公开发布；", "6. 鼓励各国交流。"]

    def _indents(self):
        result = run("working-paper", save(source(self.LINES)))
        self.assertEqual(errors(result), [])
        return {own_text(p)[:4]: indent_of(p) for p in body_xml(result.content).iter(q("p")) if own_text(p)}

    def test_letters_that_are_also_numerals_follow_their_sequence(self):
        indents = self._indents()
        self.assertEqual({indents[key]["left"] for key in ("（a）决", "（b）资", "（c）尊", "（d）定")}, {indents["（a）决"]["left"]})
        self.assertGreater(int(indents["（i）按"]["left"]), int(indents["（d）定"]["left"]), "(i) after (d) is a numeral, one level deeper")
        self.assertEqual(indents["（i）按"], indents["（ii）"])

    def test_full_width_period_marks_a_numbered_item(self):
        indents = self._indents()
        self.assertEqual(indents["5．建立"], indents["6. 鼓"])


class ParenDigitTests(unittest.TestCase):
    def test_parenthesized_digits_nest_under_the_item_that_introduces_them(self):
        lines = ["工作文件3.2", "联合国环境大会", "塑料污染", "起草国：日本国",
                 "1. 建立技术交流平台：", "（1）汇总可复制的减塑实践；", "（2）每半年提交进展说明；", "2. 鼓励自愿资金支持。"]
        result = run("working-paper", save(source(lines)))
        self.assertEqual(errors(result), [])
        indents = {own_text(p)[:3]: indent_of(p) for p in body_xml(result.content).iter(q("p")) if own_text(p)}
        self.assertEqual(indents["（1）"], indents["（2）"])
        self.assertGreater(int(indents["（1）"]["left"]), int(indents["1. "]["left"]))
        self.assertEqual(indents["1. "], indents["2. "])


def run_props(content: bytes, needle: str) -> set:
    body = body_xml(content)
    run_ = next(r for r in body.iter(q("r")) if needle in "".join(t.text or "" for t in r.iter(q("t"))))
    rpr = run_.find(q("rPr"))
    return set() if rpr is None else {child.tag.split("}")[-1] for child in rpr}


class VisibilityTests(unittest.TestCase):
    """Formatting never changes what is visible or what it means."""

    def test_a_hidden_note_stays_hidden_and_is_reported(self):
        document = source()
        paragraph = document.add_paragraph()
        for xml in ('<w:r><w:t xml:space="preserve">第一条 决定继续审议</w:t></w:r>',
                    '<w:r><w:rPr><w:vanish/></w:rPr><w:t>（内部备注：勿公开）</w:t></w:r>', '<w:r><w:t>。</w:t></w:r>'):
            paragraph._p.append(fragment(xml))
        result = run("draft-resolution", save(document))
        self.assertEqual(errors(result), [])
        self.assertIn("vanish", run_props(result.content, "内部备注"))
        self.assertTrue(any(item.code == "hidden-text" for item in result.validations))

    def test_an_amendments_strikethrough_is_kept(self):
        document = source(["非友好修正案1.3.2", "联合国大会", "修正案测试", "起草国：法兰西共和国", "附议国：日本国"])
        paragraph = document.add_paragraph()
        for xml in ('<w:r><w:t xml:space="preserve">将第二条修改为：“决定</w:t></w:r>', '<w:r><w:rPr><w:strike/></w:rPr><w:t>立即</w:t></w:r>',
                    '<w:r><w:t>继续审议此问题。”</w:t></w:r>'):
            paragraph._p.append(fragment(xml))
        result = run("unfriendly-amendment", save(document))
        self.assertEqual(errors(result), [])
        self.assertIn("strike", run_props(result.content, "立即"))


class FontPartTests(unittest.TestCase):
    """Font table, theme and settings: the same in both engines; equations keep their font."""

    def test_font_parts(self):
        result = run("draft-resolution", save(source(HEADER + ["第一条 决定继续审议。"])))
        self.assertEqual(errors(result), [])
        archive = zipfile.ZipFile(BytesIO(result.content))
        settings = archive.read("word/settings.xml").decode("utf-8")
        theme = archive.read("word/theme/theme1.xml").decode("utf-8")
        table = archive.read("word/fontTable.xml").decode("utf-8")
        self.assertIn('m:val="Cambria Math"', settings, "Times New Roman has no math table")
        self.assertNotIn("Calibri", theme)
        self.assertNotIn("Times New Roman Light", theme)
        self.assertNotIn('w:name="Calibri"', table)
        self.assertRegex(table, r'w:name="SimSun"><w:altName w:val="Songti SC"/>')


class SectionTests(unittest.TestCase):
    def test_a_source_document_grid_is_cleared_in_every_type(self):
        document = source(["友好修正案1.2", "联合国大会", "修正案测试", "起草国：法兰西共和国", "附议国：日本国", "在第一条后增加一款。"])
        document.sections[0]._sectPr.append(fragment('<w:docGrid w:type="lines" w:linePitch="312" w:charSpace="0"/>'))
        result = run("friendly-amendment", save(document))
        self.assertEqual(errors(result), [])
        self.assertEqual(list(body_xml(result.content).iter(q("docGrid"))), [])


class TableEmphasisTests(unittest.TestCase):
    def test_character_style_emphasis_in_a_table_cell_survives(self):
        document = source(["立场文件", "委员会：联合国大会", "议题：测试议题", "国家/席位：日本国", "代表：测试代表", "日本的立场如下。"])
        document.styles.element.append(fragment('<w:style w:type="character" w:styleId="Stress"><w:name w:val="Stress"/><w:rPr><w:i/></w:rPr></w:style>'))
        cell = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        cell._p.append(fragment('<w:r><w:t>表内</w:t></w:r>'))
        cell._p.append(fragment('<w:r><w:rPr><w:rStyle w:val="Stress"/></w:rPr><w:t>重点词</w:t></w:r>'))
        result = run("position-paper", save(document))
        self.assertEqual(errors(result), [])
        paragraph = next(p for p in body_xml(result.content).iter(q("p")) if own_text(p) == "表内重点词")
        self.assertEqual(italic_chars(paragraph), "重点词")


class PackageTests(unittest.TestCase):
    def test_utf16_part_with_a_dtd_is_rejected_as_unreadable(self):
        document = source(["决议草案1.0", "第一条 决定。"])
        data = save(document)
        archive = zipfile.ZipFile(BytesIO(data))
        xml = archive.read("word/document.xml").decode("utf-8")
        start = xml.index("<w:document")
        evil = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE w:document [<!ENTITY a "x">]>' + xml[start:]
        out = BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
            for item in archive.infolist():
                payload = ("﻿" + evil).encode("utf-16-le") if item.filename == "word/document.xml" else archive.read(item.filename)
                target.writestr(item.filename, payload)
        with self.assertRaises(InvalidDocxError):
            validate_docx_package(out.getvalue())


class CliTests(unittest.TestCase):
    def test_word_lock_files_in_a_folder_are_skipped_not_failed(self):
        import tempfile

        sys.path.insert(0, str(ROOT / "backend"))
        from cli import run as cli_run

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            (folder / "doc.docx").write_bytes(save(source(HEADER + ["第一条 决定继续审议。"])))
            (folder / "~$doc.docx").write_bytes(b"\x0bowner")  # what Word leaves while the file is open
            status = cli_run([str(folder), "--type", "draft-resolution", "--output-dir", str(folder / "out")])
            self.assertEqual(status, 0)
            self.assertTrue((folder / "out" / "doc_formatted.docx").is_file())


@unittest.skipUnless((ROOT / ".git").exists(), "needs the Git checkout")
class PackagingTests(unittest.TestCase):
    def test_release_preserves_git_executable_modes_even_if_working_mode_differs(self):
        import shutil
        import subprocess
        import tempfile
        with tempfile.TemporaryDirectory() as folder:
            clone = Path(folder) / "repo"
            subprocess.run(["git", "clone", "-q", str(ROOT), str(clone)], check=True)
            shutil.copy(ROOT / "scripts" / "package-release.py", clone / "scripts" / "package-release.py")
            executable = clone / "start helper"
            executable.write_text("#!/bin/sh\nexit 0\n")
            executable.chmod(0o755)
            subprocess.run(["git", "-C", str(clone), "add", "scripts", "start helper"], check=True)
            executable.chmod(0o644)
            archive_path = Path(folder) / "out.zip"
            subprocess.run([sys.executable, str(clone / "scripts" / "package-release.py"), "--output", str(archive_path)],
                           capture_output=True, text=True, check=True)
            with zipfile.ZipFile(archive_path) as archive:
                self.assertEqual((archive.getinfo("Munword/start helper").external_attr >> 16) & 0o777, 0o755)

    def test_an_untracked_document_is_never_packaged(self):
        import shutil
        import subprocess
        import tempfile

        with tempfile.TemporaryDirectory() as folder:
            clone = Path(folder) / "repo"
            subprocess.run(["git", "clone", "-q", str(ROOT), str(clone)], check=True)
            shutil.copy(ROOT / "scripts" / "package-release.py", clone / "scripts" / "package-release.py")
            shutil.copy(ROOT / "scripts" / "audit-source.py", clone / "scripts" / "audit-source.py")
            subprocess.run(["git", "-C", str(clone), "add", "scripts"], check=True)
            (clone / "我的决议.docx").write_bytes(save(source(HEADER + ["第一条 决定继续审议。"])))
            completed = subprocess.run([sys.executable, str(clone / "scripts" / "package-release.py"), "--output", str(Path(folder) / "out.zip")],
                                       capture_output=True, text=True)
            self.assertNotEqual(completed.returncode, 0)
            self.assertIn("我的决议.docx", completed.stderr)
            self.assertFalse((Path(folder) / "out.zip").exists())


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app, base_url="http://127.0.0.1:8000")
        self.content = save(source(HEADER + ["第一条 决定继续审议。"]))

    def _post(self, **fields):
        return self.client.post(
            "/api/format/draft-resolution",
            files={"file": ("a.docx", self.content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data=fields,
        )

    def test_overrides_that_are_not_an_object_are_an_invalid_request(self):
        for value in ("[]", "1", '"text"'):
            response = self._post(overrides_json=value)
            self.assertEqual(response.status_code, 422, value)

    def test_only_loopback_hosts_are_served(self):
        response = TestClient(app, base_url="http://testserver").get("/api/health")
        self.assertEqual(response.status_code, 400)


class ReleaseBoundaryTests(unittest.TestCase):
    def test_whole_document_hidden_default_is_cleared(self):
        document = source(HEADER + ["第一条 决定继续审议。"])
        document.styles["Normal"].element.get_or_add_rPr().append(fragment("<w:vanish/>"))
        result = run("draft-resolution", save(document))
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            for name in ("word/document.xml", "word/styles.xml"):
                self.assertEqual(len(list(etree.fromstring(archive.read(name)).iter(q("vanish")))), 0, name)

    def test_inherited_visibility_and_deletion_marks_survive(self):
        for prop in ("vanish", "webHidden", "specVanish", "strike", "dstrike"):
            with self.subTest(prop=prop):
                document = source()
                document.styles.element.append(fragment(
                    f'<w:style w:type="character" w:styleId="PrivateNote"><w:name w:val="PrivateNote"/><w:rPr><w:{prop}/></w:rPr></w:style>'))
                p = document.add_paragraph("第一条 决定")
                r = p.add_run("内部备注")
                r._r.get_or_add_rPr().append(fragment('<w:rStyle w:val="PrivateNote"/>'))
                p.add_run("继续审议。")
                result = run("draft-resolution", save(document))
                self.assertEqual(errors(result), [])
                marked = next(r for r in body_xml(result.content).iter(q("r")) if "内部备注" in "".join(r.itertext()))
                self.assertIsNotNone(marked.find(f'{q("rPr")}/{q(prop)}'))
                self.assertTrue(any(item.code == "hidden-text" for item in result.validations))

    def test_theme_changes_exact_font_faces_only(self):
        from app.fonts import normalize_font_parts
        data = BytesIO()
        with zipfile.ZipFile(data, "w") as archive:
            archive.writestr("word/theme/theme1.xml", '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Cambria Sample"><a:majorFont><a:latin typeface="Calibri Light"/><a:font script="Math" typeface="Cambria Math"/></a:majorFont></a:theme>')
        result = normalize_font_parts(data.getvalue(), "zh")
        xml = zipfile.ZipFile(BytesIO(result)).read("word/theme/theme1.xml").decode()
        self.assertIn('typeface="Times New Roman"', xml)
        self.assertIn('typeface="Cambria Math"', xml)
        self.assertIn('name="Cambria Sample"', xml)

    def test_bomless_utf16_dtd_with_leading_whitespace_is_rejected(self):
        from app.docx_package import declares_dtd
        text = ' \n<!DOCTYPE x [<!ENTITY a "x">]><x/>'
        for encoding in ("utf-16-le", "utf-16-be"):
            self.assertTrue(declares_dtd(text.encode(encoding)))


if __name__ == "__main__":
    unittest.main()
