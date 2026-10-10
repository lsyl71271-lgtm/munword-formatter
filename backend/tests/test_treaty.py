"""Diplomatic agreements and joint statements (mirror of tests/treaty.test.mjs)."""

import copy
import json
import sys
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
from app import content_guard  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402
from app.treaty import treaty_parties  # noqa: E402

TEMPLATES = ROOT / "templates" / "pkunmun2026"
SAMPLES = ROOT / "examples" / "acceptance-inputs"


def text(paragraph):
    parts = []
    for node in paragraph.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab") and node.getparent().tag == qn("w:r"):
            parts.append("\t")
    return "".join(parts)


def attr(paragraph, tag, name):
    node = next(paragraph.iter(qn(f"w:{tag}")), None)
    return None if node is None else node.get(qn(f"w:{name}"))


def docx(lines):
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    data = BytesIO()
    document.save(data)
    return data.getvalue()


class TreatyTests(unittest.TestCase):
    def run_type(self, content, document_type, overrides=None, **options):
        result = PIPELINES[document_type](TEMPLATES).run(content, overrides=overrides, **options)
        self.assertEqual([item.detail for item in result.validations if item.status == "error"], [])
        body = Document(BytesIO(result.content)).element.body
        return result, list(body.iter(qn("w:p")))

    @staticmethod
    def signature_rows(paragraphs):
        return [text(p) for p in paragraphs if next(p.iter(qn("w:tabs")), None) is not None]

    def test_parties_come_from_the_title(self):
        self.assertEqual(treaty_parties(["日本国、巴西联邦共和国与加拿大关于海洋合作的协定"], "zh"), ["日本国", "巴西联邦共和国", "加拿大"])
        self.assertEqual(
            treaty_parties(["中华人民共和国、法兰西共和国、大不列颠及北爱尔兰联合王国和德意志联邦共和国关于搜救的协定"], "zh"),
            ["中华人民共和国", "法兰西共和国", "大不列颠及北爱尔兰联合王国", "德意志联邦共和国"],
        )
        self.assertEqual(treaty_parties(["日本国、法兰西共和国、加拿大", "就海洋治理的联合声明"], "zh"), ["日本国", "法兰西共和国", "加拿大"])
        self.assertEqual(treaty_parties(["日本国、法兰西共和国和加拿大联合声明"], "zh"), ["日本国", "法兰西共和国", "加拿大"])
        self.assertEqual(treaty_parties(["Joint Statement of Japan, Trinidad and Tobago and Canada on Ocean Monitoring"], "en"), ["Japan", "Trinidad and Tobago", "Canada"])
        self.assertEqual(treaty_parties(["Agreement between Japan and Canada on Fisheries"], "en"), ["Japan", "Canada"])
        self.assertEqual(treaty_parties(["工作文件 [编号]"], "zh"), [])

    def test_agreement_layout_and_added_representative(self):
        result, paragraphs = self.run_type((SAMPLES / "12_中文外交协定.docx").read_bytes(), "diplomatic-agreement")
        self.assertEqual(result.model.sponsors, ["日本国", "巴西联邦共和国", "加拿大"])
        title = paragraphs[0]
        self.assertEqual([attr(title, "jc", "val"), attr(title, "spacing", "line"), attr(title, "spacing", "lineRule"), attr(title, "sz", "val"), attr(title, "b", "val")],
                         ["center", "624", "exact", "32", "1"])
        chapter = next(p for p in paragraphs if text(p) == "第一章 合作范围")
        self.assertEqual([attr(chapter, "jc", "val"), attr(chapter, "b", "val")], ["center", "1"])
        article = next(p for p in paragraphs if text(p).startswith("第一条"))
        self.assertEqual([attr(article, "jc", "val"), attr(article, "spacing", "line"), attr(article, "ind", "firstLine")], ["both", "468", "480"])
        self.assertFalse([p for p in paragraphs if not text(p).strip()])
        self.assertEqual(self.signature_rows(paragraphs), ["\t日本国代表\t巴西联邦共和国代表\t加拿大代表", "\t示例甲\t示例乙\t"])

    def test_statement_rows_numbers_and_quotation(self):
        result, paragraphs = self.run_type((SAMPLES / "13_中文联合声明.docx").read_bytes(), "joint-statement")
        texts = [text(p) for p in paragraphs]
        self.assertIn("2026年5月1日，四方在东京举行会谈，共同声明如下：", texts)
        self.assertIn("1. 四方重申对《联合国海洋法公约》的承诺。", texts)
        quote = next(p for p in paragraphs if text(p).startswith("“"))
        self.assertEqual([attr(quote, "ind", "left"), attr(quote, "i", "val")], ["240", "1"])
        self.assertEqual(self.signature_rows(paragraphs), ["\t日本国代表\t法兰西共和国代表", "\t\t", "\t巴西联邦共和国代表\t加拿大代表", "\t\t"])

    def test_english_statement_and_step_03_parties(self):
        content = (SAMPLES / "14_English_Joint_Statement.docx").read_bytes()
        _, paragraphs = self.run_type(content, "joint-statement")
        self.assertEqual(self.signature_rows(paragraphs), [
            "\tRepresentative of Japan\tRepresentative of Trinidad and Tobago", "\tSample Delegate A\t",
            "\tRepresentative of Canada", "\tSample Delegate B",
        ])
        _, edited = self.run_type(content, "joint-statement", {"sponsors": ["Japan", "Canada"]})
        self.assertEqual(self.signature_rows(edited), ["\tRepresentative of Japan\tRepresentative of Canada", "\tSample Delegate A\tSample Delegate B"])

    def test_fewer_than_two_parties_adds_no_block(self):
        result, paragraphs = self.run_type(docx(["工作文件 [编号]", "正文一段。"]), "diplomatic-agreement")
        self.assertEqual(result.model.sponsors, [])
        self.assertEqual(self.signature_rows(paragraphs), [])
        self.assertTrue(any("至少两个签署方" in item.detail for item in result.validations))

    def test_guard_refuses_dropped_name_and_changed_statement(self):
        document = Document(BytesIO(docx(["甲国与乙国关于合作的协定", "各方同意合作。", "甲国代表    乙国代表", "张三    李四"])))
        before = content_guard.Snapshot.take(document)
        _, body, labels, names = [p._p for p in document.paragraphs]
        log = {}
        for line in ("\t甲国代表\t乙国代表", "\t张三\t"):
            element = OxmlElement("w:p")
            run = OxmlElement("w:r")
            element.append(run)
            for index, part in enumerate(line.split("\t")):
                if index:
                    run.append(OxmlElement("w:tab"))
                node = OxmlElement("w:t")
                node.text = part
                run.append(node)
            labels.addprevious(element)
            log[element] = content_guard.Edit("signature", "signature", json.dumps(["甲国代表", "乙国代表"], ensure_ascii=False), True)
        for element in (labels, names):
            log[element] = content_guard.Edit("signature", "signature")
            element.getparent().remove(element)
        next(body.iter(qn("w:t"))).text = "1. 各方同意合作！"
        log[body] = content_guard.Edit("statement-number", expected="1. 各方同意合作！")
        problems = "；".join(content_guard.verify_format(before, document, log, allowed_titles=(), labels=()))
        self.assertIn("签字栏的姓名被删除、改写或新增", problems)
        self.assertIn("编号以外的内容发生变化", problems)

    def test_draft_resolution_keeps_the_uploaded_name(self):
        resolution, _ = self.run_type((SAMPLES / "07_中文决议草案.docx").read_bytes(), "draft-resolution", source_name="我的决议草案 第三稿.docx")
        self.assertEqual(resolution.filename, "我的决议草案 第三稿.docx")
        agreement, _ = self.run_type((SAMPLES / "12_中文外交协定.docx").read_bytes(), "diplomatic-agreement", source_name="随便的名字.docx")
        self.assertEqual(agreement.filename, "外交协定 日本国 v1.docx")


if __name__ == "__main__":
    unittest.main()
