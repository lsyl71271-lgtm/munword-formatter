from __future__ import annotations

import sys
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app.countries import sort_countries  # noqa: E402
from app.docx_package import validate_docx_package  # noqa: E402
from app.fingerprint import canonical_body, visible_text_signature  # noqa: E402
from app.formatters.base import ContentSnapshot  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402


class PipelineTests(unittest.TestCase):
    def source(self, paragraphs: list[str]) -> bytes:
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    def without_numbering_part(self, content: bytes) -> bytes:
        """Make a valid plain DOCX that has no numbering relationship or part."""

        source = zipfile.ZipFile(BytesIO(content))
        output = BytesIO()
        with source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                if item.filename == "word/numbering.xml":
                    continue
                payload = source.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    root = ElementTree.fromstring(payload)
                    for child in list(root):
                        if child.get("PartName") == "/word/numbering.xml":
                            root.remove(child)
                    payload = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                elif item.filename == "word/_rels/document.xml.rels":
                    root = ElementTree.fromstring(payload)
                    for child in list(root):
                        if child.get("Type", "").endswith("/numbering"):
                            root.remove(child)
                    payload = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(item, payload)
        return output.getvalue()

    def test_all_six_pipelines_produce_openable_docx(self):
        samples = {
            "position-paper": ["立场文件", "委员会：联合国大会", "国家/席位：日本国", "代表：甲", "正文保持不变。"],
            "working-paper": ["工作文件 [编号]", "委员会：联合国大会", "起草国：日本国，巴西联邦共和国", "1. 建立机制", "(a) 提交报告"],
            "draft-directive": ["指令草案 [编号]", "委员会：安全理事会", "起草国：日本国", "附议国：加拿大", "安全理事会，", "1. 要求提交报告", "2. 决定建立机制"],
            "draft-resolution": ["决议草案 [编号]", "委员会：联合国大会", "起草国：日本国", "附议国：加拿大", "联合国大会，", "回顾相关决议，", "第一条 要求提交报告；", "第二条 决定建立机制。"],
            "friendly-amendment": ["友好修正案 [编号]", "委员会：联合国大会", "起草国：日本国", "附议国：加拿大", "1. 加入限定语。"],
            "unfriendly-amendment": ["非友好修正案 [编号]", "委员会：联合国大会", "起草国：法国", "附议国：加拿大", "1. 删除第三条。"],
        }
        for document_type, paragraphs in samples.items():
            with self.subTest(document_type=document_type):
                source = self.source(paragraphs)
                result = PIPELINES[document_type](ROOT / "templates" / "pkunmun2026").run(source)
                reopened = Document(BytesIO(result.content))
                self.assertTrue(reopened.paragraphs)
                self.assertFalse(any(item.status == "error" for item in result.validations), result.validations)
                self.assertEqual(canonical_body(Document(BytesIO(source))), canonical_body(reopened))

    def test_position_repairs_missing_structure_without_sample_matching(self):
        source = self.source([
            "立场文件7.9", "：虚构委员会", "全新且未出现过的议题", "测试共和国", "新代表",
            "导言是一段足够长的原创正文，用来确认规则不会依赖任何已保存的句子。",
            "第一部分概述：", "新的安全议题",
            "这是该节的长段正文，用来让系统可靠判断上一段是短标题而不是普通正文。",
            "（二）新的经济议题",
        ])
        pipeline = PIPELINES["position-paper"](ROOT / "templates" / "pkunmun2026")
        model = pipeline.parse(source)
        self.assertEqual((model.committee, model.topic, model.country), ("虚构委员会", "全新且未出现过的议题", "测试共和国"))
        output = [paragraph.text for paragraph in Document(BytesIO(pipeline.run(source).content)).paragraphs]
        self.assertEqual(output[1:5], ["委员会：虚构委员会", "议题：全新且未出现过的议题", "国家：测试共和国", "代表：新代表"])
        self.assertIn("（一）新的安全议题", output)

    def test_resolution_repairs_embedded_markers_generically(self):
        source = self.source([
            "决议草案9.8", "虚构委员会", "原创议题", "起草国：测试共和国", "附议国：示例联邦",
            "虚构委员会，", "回顾原创原则，", "第一条 确定通用规则：（一）保留一级分款换行；",
            "第二条 建立机制：（子）补充分支；", "第三条 继续执行：（子子）更深分支。",
        ])
        result = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026").run(source)
        output = Document(BytesIO(result.content))
        texts = [paragraph.text for paragraph in output.paragraphs]
        self.assertIn("第一条 确定通用规则：\n（一）保留一级分款换行；", texts)
        self.assertIn("（子）补充分支；", texts)
        self.assertIn("（子子）更深分支。", texts)
        self.assertEqual(
            "".join(paragraph.text for paragraph in Document(BytesIO(source)).paragraphs).replace("\n", ""),
            "".join(texts).replace("\n", ""),
        )

    def test_resolution_preserves_deliberate_tab_separated_markers(self):
        source = self.source([
            "决议草案9.9", "虚构委员会", "原创议题", "起草国：测试共和国", "附议国：示例联邦",
            "虚构委员会，", "回顾原创原则，", "第一条 决定建立机制：\t（一）保留同段分款；",
        ])
        result = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026").run(
            source, normalize_punctuation=False
        )
        self.assertEqual(result.model.repair_actions, [])
        self.assertEqual(
            visible_text_signature(Document(BytesIO(source))),
            visible_text_signature(Document(BytesIO(result.content))),
        )

    def test_resolution_removes_extreme_formatting_noise(self):
        source = self.source([
            "决议草案8.8", "虚构委员会", "原创议题", "起草国：测试共和国", "附议国：示例联邦",
            "虚构委员会，", "回顾原创原则，", "第一条 决定建立机制。",
        ])
        damaged = Document(BytesIO(source))
        section = damaged.sections[0]
        width, height = section.page_width, section.page_height
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = max(width, height), min(width, height)
        section.left_margin = section.right_margin = Inches(4)
        grid = OxmlElement("w:docGrid")
        section._sectPr.append(grid)
        paragraph = damaged.paragraphs[-1]
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.left_indent = Inches(5)
        shading = OxmlElement("w:shd")
        paragraph._p.get_or_add_pPr().append(shading)
        for run in paragraph.runs:
            run.font.size = Pt(28)
            rpr = run._r.get_or_add_rPr()
            rpr.append(OxmlElement("w:vanish"))
            rpr.append(OxmlElement("w:strike"))
        stream = BytesIO()
        damaged.save(stream)
        result = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026").run(
            stream.getvalue(), normalize_punctuation=False
        )
        repaired = Document(BytesIO(result.content))
        # The only text change is the handbook's unlabeled committee / topic lines (页31).
        expected = visible_text_signature(damaged).replace("委员会:", "", 1).replace("议题:", "", 1)
        self.assertEqual(expected, visible_text_signature(repaired))
        self.assertGreater(repaired.sections[0].page_height, repaired.sections[0].page_width)
        self.assertFalse(repaired.paragraphs[-1].paragraph_format.page_break_before)
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        for tag in ("w:docGrid", "w:shd"):
            self.assertNotIn(f"<{tag}", xml)
        # Hidden or struck text that covers only part of the document is the
        # author's (a private note, an amendment's deletion): it is kept as it
        # was and reported, never shown.  Whole-document hiding is damage.
        self.assertIn("<w:vanish", xml)
        self.assertTrue(any(item.code == "hidden-text" for item in result.validations))

    def test_docx_package_validation_rejects_plain_bytes(self):
        with self.assertRaisesRegex(ValueError, "有效的 DOCX"):
            validate_docx_package(b"not a docx")

    def test_step03_rewrites_only_changed_field(self):
        source = self.source(["立场文件", "委员会：原委员会", "议题：原议题", "国家：原国家", "代表：原代表", "正文绝对不能变化。"])
        result = PIPELINES["position-paper"](ROOT / "templates" / "pkunmun2026").run(source, overrides={"committee": "新委员会"})
        output = [paragraph.text for paragraph in Document(BytesIO(result.content)).paragraphs]
        for expected in ("委员会：新委员会", "议题：原议题", "国家：原国家", "代表：原代表", "正文绝对不能变化。"):
            self.assertIn(expected, output)

    def test_country_parser_stops_before_unfriendly_amendment_body(self):
        source = self.source([
            "非友好修正案1.3.2",
            "起草国：德意志联邦共和国",
            "附议国：伊朗伊斯兰共和国、巴西联邦共和国",
            "对决议草案第一条补充：成员国捐款为1500万。",
        ])
        pipeline = PIPELINES["unfriendly-amendment"](ROOT / "templates" / "pkunmun2026")
        model = pipeline.parse(source)
        self.assertEqual(model.signatories, ["伊朗伊斯兰共和国", "巴西联邦共和国"])
        self.assertEqual([clause.text for clause in model.operative_clauses], ["对决议草案第一条补充：成员国捐款为1500万。"])

    def test_working_paper_preserves_visible_numbering(self):
        source = self.source(["工作文件 [编号]", "委员会：联合国大会", "起草国：日本国", "1. 建立机制", "(a) 提交报告"])
        result = PIPELINES["working-paper"](ROOT / "templates" / "pkunmun2026").run(source)
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            numbering = archive.read("word/numbering.xml").decode("utf-8")
            document = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("1. 建立机制", document)
        self.assertIn("(a) 提交报告", document)
        reopened = Document(BytesIO(result.content))
        self.assertTrue(all(run.font.size and run.font.size.pt == 12 for paragraph in reopened.paragraphs for run in paragraph.runs if run.text))

    def test_working_paper_removes_extreme_noise_without_losing_numbering(self):
        source = self.source([
            "工作文件8.8",
            "委员会：联合国教科文组织",
            "议题：原创议题",
            "起草国：测试共和国",
            "1. 建立恢复机制，",
            "a）提交季度报告，",
            "下一会期继续讨论。",
        ])
        damaged = Document(BytesIO(source))
        damaged.styles["Normal"].font.size = Pt(14)
        section = damaged.sections[0]
        width, height = section.page_width, section.page_height
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = max(width, height), min(width, height)
        section.left_margin = section.right_margin = Inches(4)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "3")
        section._sectPr.append(cols)
        original_num_ids = []
        for index, paragraph in enumerate(damaged.paragraphs):
            paragraph.paragraph_format.page_break_before = index > 0
            paragraph.paragraph_format.left_indent = Inches(5)
            ppr = paragraph._p.get_or_add_pPr()
            ppr.append(OxmlElement("w:shd"))
            for run_index, run in enumerate(paragraph.runs):
                run.font.size = Pt(6 if run_index % 2 else 28)
                rpr = run._r.get_or_add_rPr()
                for tag in ("vanish", "strike", "smallCaps", "color", "highlight"):
                    rpr.append(OxmlElement(f"w:{tag}"))
        for paragraph_index, num_id in ((4, 1), (6, 2)):
            ppr = damaged.paragraphs[paragraph_index]._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            num = OxmlElement("w:numId")
            num.set(qn("w:val"), str(num_id))
            num_pr.extend((ilvl, num))
            ppr.append(num_pr)
            original_num_ids.append(str(num_id))
        stream = BytesIO()
        damaged.save(stream)

        result = PIPELINES["working-paper"](ROOT / "templates" / "pkunmun2026").run(
            stream.getvalue(), normalize_punctuation=False
        )
        repaired = Document(BytesIO(result.content))
        # The only text change is the handbook's unlabeled committee / topic lines (页31).
        expected = visible_text_signature(damaged).replace("委员会:", "", 1).replace("议题:", "", 1)
        self.assertEqual(expected, visible_text_signature(repaired))
        self.assertAlmostEqual(repaired.sections[0].page_width.mm, 210, places=1)
        self.assertAlmostEqual(repaired.sections[0].left_margin.mm, 31.75, places=1)
        self.assertTrue(all(run.font.size and run.font.size.pt == 12 for paragraph in repaired.paragraphs for run in paragraph.runs if run.text))
        self.assertTrue(all(run.bold is not True and run.italic is not True and not run.underline for paragraph in repaired.paragraphs[4:] for run in paragraph.runs if run.text))
        self.assertTrue(all(run.bold for run in repaired.paragraphs[0].runs if run.text))
        sponsor_runs = [run for run in repaired.paragraphs[3].runs if run.text and "测试共和国" in run.text]
        self.assertTrue(sponsor_runs and all(run.bold and run.italic for run in sponsor_runs))
        repaired_num_ids = [
            str(repaired.paragraphs[index]._p.pPr.numPr.numId.val)
            for index in (4, 6)
        ]
        self.assertEqual(repaired_num_ids, original_num_ids)
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        for tag in ("w:cols", "w:vanish", "w:strike", "w:smallCaps", "w:color", "w:highlight", "w:shd"):
            self.assertNotIn(f"<{tag}", xml)

    def test_missing_numbering_part_stays_optional_for_visible_markers(self):
        source = self.without_numbering_part(self.source([
            "工作文件 [编号]",
            "委员会：联合国大会",
            "起草国：日本国",
            "1. 建立合作机制",
            "2. 定期提交报告",
        ]))
        with zipfile.ZipFile(BytesIO(source)) as archive:
            self.assertNotIn("word/numbering.xml", archive.namelist())

        result = PIPELINES["working-paper"](ROOT / "templates" / "pkunmun2026").run(source)
        reopened = Document(BytesIO(result.content))
        self.assertTrue(reopened.paragraphs)
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            self.assertNotIn("word/numbering.xml", archive.namelist())
        self.assertIn("1. 建立合作机制", [paragraph.text for paragraph in reopened.paragraphs])

    def test_chinese_position_paper_uses_chinese_article_markers(self):
        source = self.source([
            "立场文件",
            "委员会：欧洲联盟",
            "议题：测试议题",
            "国家/席位：瑞典王国",
            "代表：甲",
            "（一）俄格冲突",
            "1、要求建立核查机制",
        ])
        result = PIPELINES["position-paper"](ROOT / "templates" / "pkunmun2026").run(source)
        reopened = Document(BytesIO(result.content))
        # Handbook 页15–16: one blank line after the header, proposals start
        # two characters in (24 pt) with their text at 45 pt.
        self.assertEqual(reopened.paragraphs[5].text, "")
        self.assertEqual(reopened.paragraphs[6].text, "（一）俄格冲突")
        proposal = reopened.paragraphs[7]
        self.assertEqual(proposal.text, "1、\t要求建立核查机制")
        self.assertIsNone(proposal._p.pPr.numPr)
        self.assertEqual(proposal.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertAlmostEqual(proposal.paragraph_format.left_indent.pt, 45, places=1)
        self.assertAlmostEqual(proposal.paragraph_format.first_line_indent.pt, -21, places=1)
        self.assertTrue(all(run.font.name == "Times New Roman" for run in proposal.runs if run.text))
        statuses = {item.code: item.status for item in result.validations}
        self.assertEqual(statuses["position-metadata"], "pass")
        self.assertEqual(statuses["position-section-markers"], "pass")
        self.assertEqual(statuses["position-font-size"], "pass")

    def test_position_validation_rejects_missing_labels_and_section_markers(self):
        source = self.source([
            "立场文件",
            "委员会：欧洲联盟",
            "议题：测试议题",
            "国家：瑞典王国",
            "代表：甲",
            "（一）俄格冲突",
            "正文保持不变。",
        ])
        pipeline = PIPELINES["position-paper"](ROOT / "templates" / "pkunmun2026")
        model = pipeline.parse(source)
        broken = Document(BytesIO(source))
        broken.paragraphs[5].text = "俄格冲突"
        before = ContentSnapshot.take(Document(BytesIO(source)))
        after = ContentSnapshot.take(broken)
        # The generic content fingerprint intentionally ignores labels and
        # list markers, so position papers need dedicated structural gates.
        self.assertEqual(before.canonical, after.canonical)
        # A real run normalizes punctuation by default, which lets the
        # canonical fingerprint stand in for the strict visible signature.
        pipeline.formatter._normalize_punctuation_requested = True
        validations = pipeline.formatter._validate(broken, model, before, after, False)
        statuses = {item.code: item.status for item in validations}
        self.assertEqual(statuses["content"], "pass")
        self.assertEqual(statuses["position-metadata"], "pass")
        self.assertEqual(statuses["position-section-markers"], "error")

        broken.paragraphs[1].text = "：欧洲联盟"
        validations = pipeline.formatter._validate(broken, model, before, ContentSnapshot.take(broken), False)
        statuses = {item.code: item.status for item in validations}
        self.assertEqual(statuses["position-metadata"], "error")

    def test_position_paper_reference_block_matches_sample(self):
        source = self.source([
            "立场文件",
            "委员会：欧洲联盟",
            "议题：测试议题",
            "国家/席位：瑞典王国",
            "代表：甲",
            "正文保持十二磅。",
            "[1] European Council, Reference Entry",
            "https://example.org/reference",
        ])
        result = PIPELINES["position-paper"](ROOT / "templates" / "pkunmun2026").run(source)
        reopened = Document(BytesIO(result.content))
        references = reopened.paragraphs[-2:]
        # Set like the handbook's notes (页16, 页18): 9 pt, flush left.
        for paragraph in references:
            self.assertAlmostEqual(paragraph.paragraph_format.left_indent.pt, 0, places=1)
            self.assertAlmostEqual(paragraph.paragraph_format.first_line_indent.pt, 0, places=1)
            self.assertAlmostEqual(paragraph.paragraph_format.space_before.pt, 0, places=1)
            self.assertTrue(all(run.font.size and run.font.size.pt == 9 for run in paragraph.runs))

    def test_country_sorting(self):
        self.assertEqual(sort_countries(["日本国", "中国", "巴西联邦共和国"], "zh"), ["巴西联邦共和国", "日本国", "中国"])
        self.assertEqual(
            sort_countries(["埃塞俄比亚联邦民主共和国", "厄立特里亚国", "阿拉伯埃及共和国"], "zh"),
            ["阿拉伯埃及共和国", "埃塞俄比亚联邦民主共和国", "厄立特里亚国"],
        )
        self.assertEqual(sort_countries(["Japan", "Brazil", "Canada"], "en"), ["Brazil", "Canada", "Japan"])

    def test_chinese_resolution_supports_four_numbering_levels(self):
        source = self.source(["决议草案 [编号]", "委员会：联合国大会", "起草国：日本国", "附议国：加拿大", "联合国大会，", "第一条 要求建立机制；", "（一）鼓励提交报告；", "（子）建议明确方法；", "（甲）要求提供说明。"])
        result = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026").run(source)
        reopened = Document(BytesIO(result.content))
        visible = [paragraph.text for paragraph in reopened.paragraphs]
        expected = ["第一条 ", "（一）", "（子）", "（甲）"]
        clause_paragraphs = visible[-4:]
        for prefix, text in zip(expected, clause_paragraphs):
            self.assertTrue(text.startswith(prefix), (prefix, text))
        # Handbook 页42–43: "第一条" flush left, then (left, hanging)
        # 48/36, 67.5/36 and 87/36 pt.
        geometry = [
            (paragraph.paragraph_format.left_indent.pt, paragraph.paragraph_format.first_line_indent.pt)
            for paragraph in reopened.paragraphs[-4:]
        ]
        self.assertEqual(geometry, [(0, 0), (48, -36), (67.5, -36), (87, -36)])
        self.assertTrue(clause_paragraphs[-1].endswith("。"))

    def test_page_and_fonts(self):
        source = self.source(["立场文件", "委员会：联合国大会", "国家/席位：日本国", "代表：甲", "正文 ABC 123。"])
        result = PIPELINES["position-paper"](ROOT / "templates" / "pkunmun2026").run(source)
        reopened = Document(BytesIO(result.content))
        section = reopened.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, places=1)
        self.assertAlmostEqual(section.page_height.mm, 297, places=1)
        self.assertAlmostEqual(section.left_margin.mm, 31.75, places=1)
        self.assertAlmostEqual(section.top_margin.mm, 25.4, places=1)
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
            styles = archive.read("word/styles.xml").decode("utf-8")
        self.assertIn("SimSun", styles)
        self.assertNotIn("Arial", styles)
        self.assertNotIn("Calibri", styles)

    def test_resolution_uses_explicit_first_article_as_clause_boundary(self):
        source = self.source([
            "决议草案5.1",
            "非盟和平与安全理事会",
            "2023吉达会议",
            "",
            "起草国：阿尔及利亚民主人民共和国、南非共和国、尼日利亚共和国",
            "附议国：阿拉伯埃及共和国、埃塞俄比亚联邦民主共和国、厄立特里亚国、",
            "",
            "吉布提共和国、肯尼亚共和国",
            "",
            "非盟和平与安全理事会，",
            "回顾相关原则，",
            "呼吁国际社会兑现承诺，",
            "支持有关组织开展合作，",
            "",
            "第一条重申短期停火承诺：",
            "（一）要求双方遵守协议；",
            "（二）停止主动攻击；",
            "第二条支持建立沟通机制：",
            "（三）指定联络官；",
            "（子）说明紧急事件；",
            "第三条谴责针对平民的暴力行为。",
        ])
        pipeline = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026")
        model = pipeline.parse(source)
        self.assertEqual(model.committee, "非盟和平与安全理事会")
        self.assertEqual(model.topic, "2023吉达会议")
        self.assertEqual(len(model.sponsors), 3)
        self.assertEqual(len(model.signatories), 5)
        self.assertEqual([clause.text for clause in model.preambulatory_clauses][-2:], ["呼吁国际社会兑现承诺，", "支持有关组织开展合作，"])
        self.assertEqual(model.operative_clauses[0].text, "重申短期停火承诺：")
        self.assertEqual([clause.level for clause in model.operative_clauses[:6]], [0, 1, 1, 0, 1, 2])

        result = pipeline.run(source)
        self.assertFalse(any(item.status == "error" for item in result.validations), result.validations)
        reopened = Document(BytesIO(result.content))
        visible = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
        self.assertNotIn("：；", visible)
        self.assertIn("呼吁国际社会兑现承诺，", visible)
        self.assertIn("第一条重申", visible)
        self.assertIn("（一）要求双方遵守协议；", visible)
        # An item that introduces subclauses ends with a colon (页42–43).
        self.assertIn("（三）指定联络官：", visible)
        self.assertIn("（子）说明紧急事件；", visible)
        self.assertEqual(len(reopened.tables), 0)
        country_paragraphs = [paragraph for paragraph in reopened.paragraphs if paragraph.text.startswith(("起草国：", "附议国："))]
        self.assertEqual(len(country_paragraphs), 2)
        self.assertNotIn("_", "\n".join(paragraph.text for paragraph in country_paragraphs))
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertEqual(document_xml.count("<w:numPr>"), 0)
        self.assertNotIn('<w:sz w:val="21"', document_xml)
        self.assertIn('<w:sz w:val="24"', document_xml)

    def test_resolution_infers_omitted_first_article_from_following_subclauses(self):
        source = self.source([
            "决议草案1.2",
            "联合国教科文组织",
            "测试议题",
            "起草国：日本国",
            "附议国：加拿大",
            "联合国教科文组织，",
            "回顾相关原则，",
            "呼吁建立国际教育基金：",
            "（一）明确资金来源；",
            "（二）提交审计报告；",
            "第二条 鼓励加强教师培训。",
        ])
        pipeline = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026")
        model = pipeline.parse(source)
        self.assertEqual([clause.text for clause in model.preambulatory_clauses], ["回顾相关原则，"])
        self.assertEqual(model.operative_clauses[0].text, "呼吁建立国际教育基金：")
        self.assertEqual([clause.level for clause in model.operative_clauses], [0, 1, 1, 0])

        result = pipeline.run(source)
        reopened = Document(BytesIO(result.content))
        visible = [paragraph.text for paragraph in reopened.paragraphs if paragraph.text]
        self.assertIn("呼吁建立国际教育基金：", visible)
        self.assertIn("第二条 鼓励加强教师培训。", visible)

    def test_reference_layout_profiles_are_applied_per_document_type(self):
        working = PIPELINES["working-paper"](ROOT / "templates" / "pkunmun2026").run(self.source([
            "工作文件1.1",
            "委员会：联合国教科文组织",
            "起草国：日本国",
            "1. 建立恢复机制",
            "(a) 提交报告",
        ]))
        working_doc = Document(BytesIO(working.content))
        working_section = working_doc.sections[0]
        # Every document type uses the handbook page and 12 pt body.
        self.assertAlmostEqual(working_section.page_width.mm, 210, places=1)
        self.assertAlmostEqual(working_section.left_margin.mm, 31.75, places=1)
        self.assertAlmostEqual(working_doc.styles["Normal"].font.size.pt, 12, places=1)

        friendly = PIPELINES["friendly-amendment"](ROOT / "templates" / "pkunmun2026").run(self.source([
            "修正案1.2.1",
            "委员会：联合国教科文组织",
            "起草国：日本国",
            "附议国：加拿大",
            "1. 加入第三条：\"提交季度报告\"；",
        ]))
        friendly_doc = Document(BytesIO(friendly.content))
        self.assertEqual(len(friendly_doc.tables), 0)
        operation = next(paragraph for paragraph in friendly_doc.paragraphs if "加入第三条" in paragraph.text)
        # Typed numbers stay as written; only the operation verb is italic (页52).
        self.assertTrue(operation.text.startswith("1. 加入第三条"))
        self.assertEqual([run.text for run in operation.runs if run.italic], ["加入"])
        self.assertTrue(all(run.font.size.pt == 12 for run in operation.runs if run.text))

        unfriendly = PIPELINES["unfriendly-amendment"](ROOT / "templates" / "pkunmun2026").run(self.source([
            "非友好修正案1.3.2",
            "起草国：德意志联邦共和国",
            "附议国：加拿大",
            "对第一条补充具体出资金额。",
        ]))
        unfriendly_doc = Document(BytesIO(unfriendly.content))
        # Header at body size: title bold, country names bold italic (页52).
        title = unfriendly_doc.paragraphs[0]
        self.assertTrue(all(run.bold and not run.italic and run.font.size.pt == 12 for run in title.runs))
        sponsors = next(paragraph for paragraph in unfriendly_doc.paragraphs if paragraph.text.startswith("起草国"))
        self.assertTrue(all(run.bold and run.font.size.pt == 12 for run in sponsors.runs))
        self.assertTrue(sponsors.runs[-1].italic)

        resolution = PIPELINES["draft-resolution"](ROOT / "templates" / "pkunmun2026").run(self.source([
            "决议草案1.2",
            "委员会：联合国教科文组织",
            "起草国：日本国",
            "附议国：加拿大",
            "联合国教科文组织，",
            "回顾相关原则，",
            "第一条 呼吁建立机制。",
        ]))
        resolution_doc = Document(BytesIO(resolution.content))
        self.assertAlmostEqual(resolution_doc.sections[0].page_width.mm, 210, places=1)
        body_run = next(run for paragraph in resolution_doc.paragraphs for run in paragraph.runs if run.text)
        self.assertAlmostEqual(body_run.font.size.pt, 12, places=1)
        self.assertAlmostEqual(resolution_doc.paragraphs[0].paragraph_format.line_spacing.pt, 15.5, places=2)


if __name__ == "__main__":
    unittest.main()
