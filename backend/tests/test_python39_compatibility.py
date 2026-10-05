"""Exercise title/header parsing with Python 3.9's keyword-free zip contract."""

import ast
import builtins
import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app.main import app
from app.parser import DocxParser
from app.semantic_policy import DOCUMENT_PROFILES


def python39_zip(*iterables):
    # No **kwargs: exercise the old interpreter contract even on Python 3.12.
    return builtins.zip(*iterables)


def source(lines):
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


class Python39CompatibilityTests(unittest.TestCase):
    def test_backend_does_not_pass_keywords_to_builtin_zip(self):
        for path in sorted((ROOT / "backend" / "app").rglob("*.py")):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == "zip":
                    with self.subTest(file=str(path.relative_to(ROOT)), line=node.lineno):
                        self.assertEqual(node.keywords, [], "Python 3.9's zip accepts no keyword arguments")

    def test_all_document_titles_with_complete_partial_and_empty_headers(self):
        titles = {
            "position-paper": "立场文件",
            "working-paper": "工作文件1.1",
            "draft-directive": "指令草案1.1",
            "draft-resolution": "决议草案1.1",
            "friendly-amendment": "友好修正案1.1.1",
            "unfriendly-amendment": "非友好修正案1.1.1",
        }
        values = {"committee": "联合国安全理事会", "topic": "和平与安全", "country": "瑞典王国", "delegate": "测试代表"}
        with patch("app.parser.zip", python39_zip, create=True):
            for document_type, title in titles.items():
                header = [values[key] for key in DOCUMENT_PROFILES[document_type]["unlabeledHeaderFields"]]
                for count in (len(header), 1, 0):
                    with self.subTest(document_type=document_type, header_count=count):
                        content = source([title, *header[:count], "第一条 决定支持和平合作。"])
                        model = DocxParser().parse(content, document_type)
                        self.assertEqual(model.title, title)
                        self.assertIn("第一条 决定支持和平合作。", model.paragraphs)
                        if count == len(header):
                            self.assertEqual(model.committee, values["committee"])
                            self.assertEqual(model.topic, values["topic"])

    def test_resolution_upload_and_download_without_manual_overrides(self):
        content = source([
            "决议草案1.1", "联合国安全理事会", "和平与安全",
            "起草国：瑞典王国", "附议国：法兰西共和国", "联合国安全理事会，",
            "回顾联合国宪章，", "第一条 决定支持和平合作。",
        ])
        client = TestClient(app, base_url="http://127.0.0.1:8000")
        files = {"file": ("resolution.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        with patch("app.parser.zip", python39_zip, create=True):
            parsed = client.post("/api/parse/draft-resolution", files=files)
            self.assertEqual(parsed.status_code, 200, parsed.text)
            self.assertEqual(parsed.json()["committee"], "联合国安全理事会")
            formatted = client.post("/api/format/draft-resolution", files=files)
            self.assertEqual(formatted.status_code, 200, formatted.text[:500] if formatted.status_code != 200 else "")
        output = Document(BytesIO(formatted.content))
        self.assertTrue(any("决定支持和平合作" in p.text for p in output.paragraphs))
