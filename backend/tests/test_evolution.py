"""Rules learned from comparing output with the user's correct documents.

Each test builds its own small document, so none depends on the private
reference files; scripts/reference_regression.py measures those.
"""

import re
import sys
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app.docx_view import visible_text  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402

TEMPLATES = ROOT / "templates" / "pkunmun2026"
RESOLUTION_HEADER = ["决议草案1.0", "联合国教科文组织", "测试议题", "起草国：法兰西共和国", "附议国：日本国", "联合国教科文组织，"]


def _save(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _without_numbering_part(content: bytes) -> bytes:
    source = zipfile.ZipFile(BytesIO(content))
    output = BytesIO()
    with source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename == "word/numbering.xml":
                continue
            data = source.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = re.sub(rb"<Relationship [^>]*numbering\.xml\"[^>]*/>", b"", data)
            elif item.filename == "[Content_Types].xml":
                data = re.sub(rb"<Override [^>]*numbering\.xml\"[^>]*/>", b"", data)
            target.writestr(item, data)
    return output.getvalue()


def _run(document_type, content, **kwargs):
    return PIPELINES[document_type](TEMPLATES).run(content, **kwargs)


def _errors(result):
    return [item.code for item in result.validations if item.status == "error"]


def _output(result):
    return Document(BytesIO(result.content))


def _paragraph(result, prefix):
    return next(p for p in _output(result).paragraphs if visible_text(p).strip().startswith(prefix))


def _numbered_document():
    """A document with one automatic decimal list (numId 1)."""

    document = Document()
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "0")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1.")):
        node = OxmlElement(tag)
        node.set(qn("w:val"), value)
        level.append(node)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "1")
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), "0")
    num.append(ref)
    numbering.append(num)
    return document


def _number(paragraph, num_id=1, level=0):
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = level
    num_pr.get_or_add_numId().val = num_id


class RobustnessTests(unittest.TestCase):
    def test_numbering_reference_without_numbering_part_does_not_crash(self):
        document = Document()
        for text in ("友好修正案1.0", "起草国：法兰西共和国", "附议国：日本国", "对决议草案第一条："):
            document.add_paragraph(text)
        _number(document.paragraphs[-1])
        # Drop numbering.xml the way some editors export: numPr stays.
        content = _without_numbering_part(_save(document))
        with zipfile.ZipFile(BytesIO(content)) as archive:
            self.assertNotIn("word/numbering.xml", archive.namelist())
        result = _run("friendly-amendment", content)
        self.assertEqual(_errors(result), [])

    def test_sponsors_line_with_a_hyperlink_is_styled_not_blocked(self):
        document = Document()
        for text in RESOLUTION_HEADER[:3]:
            document.add_paragraph(text)
        sponsors = document.add_paragraph("起草国：")
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), document.part.relate_to("https://example.org", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "法兰西共和国"
        run.append(text)
        link.append(run)
        sponsors._p.append(link)
        for line in RESOLUTION_HEADER[4:] + ["回顾以往决议，", "第一条 决定设立观察团；"]:
            document.add_paragraph(line)
        result = _run("draft-resolution", _save(document))
        self.assertEqual(_errors(result), [])

    def test_tracked_insertion_text_gets_the_house_size(self):
        document = Document()
        for text in RESOLUTION_HEADER + ["回顾以往决议，"]:
            document.add_paragraph(text).runs[0].font.size = Pt(14)
        clause = document.add_paragraph("第一条 决定")
        clause.runs[0].font.size = Pt(14)
        insertion = OxmlElement("w:ins")
        insertion.set(qn("w:id"), "9")
        insertion.set(qn("w:author"), "A")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "设立观察团；"
        run.append(text)
        insertion.append(run)
        clause._p.append(insertion)
        result = _run("draft-resolution", _save(document))
        inserted = _paragraph(result, "第一条")._p.find(f"{qn('w:ins')}/{qn('w:r')}")
        self.assertEqual(inserted.find(f"{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val")), "24")
        self.assertEqual(_errors(result), [])

    def test_emphasis_dots_and_squeezed_text_are_removed(self):
        document = Document()
        for text in ("立场文件", "委员会：测试", "议题：测试", "国家：法兰西共和国", "代表：甲", "正文内容足够长，用于测试噪声清理的效果与稳定性。"):
            paragraph = document.add_paragraph(text)
            rpr = paragraph.runs[0]._r.get_or_add_rPr()
            dots = OxmlElement("w:em")
            dots.set(qn("w:val"), "dot")
            rpr.append(dots)
        result = _run("position-paper", _save(document))
        self.assertNotIn(b"<w:em ", result.content.replace(b"<w:em/>", b"<w:em "))


class StructureRepairTests(unittest.TestCase):
    def test_unnumbered_introduction_keeps_original_list_and_cross_reference(self):
        document = _numbered_document()
        for text in ("工作文件1.1", "委员会：联合国大会", "议题：文化保护", "起草国：中国", "提出以下建议：", "以下若干建议供进一步讨论；"):
            document.add_paragraph(text)
        for text in ("鼓励合作；", "支持上述第1项建议。"):
            _number(document.add_paragraph(text))
        result = _run("working-paper", _save(document))
        self.assertEqual(_errors(result), [])
        self.assertIsNone(_paragraph(result, "以下若干")._p.pPr.numPr)
        self.assertEqual(_paragraph(result, "鼓励合作")._p.pPr.numPr.numId.val, 1)
        self.assertEqual(visible_text(_paragraph(result, "支持上述")), "支持上述第1项建议。")

    def test_possible_missing_first_item_never_shifts_original_numbers(self):
        document = _numbered_document()
        for text in RESOLUTION_HEADER + ["回顾以往决议，", "第一条 决定下列措施："]:
            document.add_paragraph(text)
        dropped = document.add_paragraph("建立数据库；")
        for text in ("加强合作；", "定期报告；"):
            _number(document.add_paragraph(text))
        # Backspace keeps the text where the number's text began.
        dropped.paragraph_format.left_indent = Pt(0)
        for paragraph in document.paragraphs[-2:]:
            paragraph.paragraph_format.left_indent = Pt(0)
        result = _run("draft-resolution", _save(document))
        self.assertIsNone(_paragraph(result, "建立数据库")._p.pPr.numPr)
        self.assertTrue(any("保留原条号" in item for item in result.model.warnings))

    def test_flattened_subclause_is_split_keeping_author_formatting(self):
        document = Document()
        for text in RESOLUTION_HEADER + ["回顾以往决议，"]:
            document.add_paragraph(text)
        article = document.add_paragraph("第一条 ")
        verb = article.add_run("决定")
        verb.italic = True
        article.add_run("下列事项：（子）设立基金；")
        result = _run("draft-resolution", _save(document))
        first = _paragraph(result, "第一条")
        self.assertEqual(visible_text(first).strip(), "第一条 决定下列事项：")
        self.assertTrue(any(run.italic and run.text == "决定" for run in first.runs))
        # The last clause of the resolution ends with a full stop (页43).
        self.assertEqual(visible_text(_paragraph(result, "（子）")).strip(), "（子）设立基金。")

    def test_numbering_kept_in_a_paragraph_style_survives(self):
        document = _numbered_document()
        style = document.styles.add_style("Numbered Clause", 1)
        num_pr = style.element.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = 1
        for text in ("工作文件1.0", "委员会：联合国教科文组织", "议题：测试", "起草国：法兰西共和国"):
            document.add_paragraph(text)
        document.add_paragraph("建立信息共享机制；", style="Numbered Clause")
        result = _run("working-paper", _save(document))
        self.assertIsNotNone(_paragraph(result, "建立信息共享机制")._p.pPr.numPr)

    def test_numid_zero_is_not_list_membership(self):
        document = Document()
        for text in RESOLUTION_HEADER + ["回顾以往决议，", "第一条 决定下列事项："]:
            document.add_paragraph(text)
        for text in ("（一）设立基金；", "（二）定期报告；", "（三）加强合作；"):
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.first_line_indent = Pt(21)
        # WPS writes numId 0 ("numbering off"); the indentation got lost.
        damaged = document.paragraphs[-1]
        _number(damaged, num_id=0)
        damaged.paragraph_format.first_line_indent = Pt(0)
        result = _run("draft-resolution", _save(document))
        # All three subclauses are list items at the handbook's second level.
        for marker in ("（一）", "（二）", "（三）"):
            fmt = _paragraph(result, marker).paragraph_format
            self.assertEqual((fmt.left_indent.pt, fmt.first_line_indent.pt), (48, -36), marker)


class ReferenceStyleTests(unittest.TestCase):
    def test_chinese_resolution_header_lines_are_bold_and_verbs_of_colon_articles_italic(self):
        document = Document()
        for text in RESOLUTION_HEADER + ["回顾以往决议，", "第一条 决定下列事项：", "（一）设立基金；"]:
            document.add_paragraph(text)
        result = _run("draft-resolution", _save(document))
        committee = _paragraph(result, "联合国教科文组织")
        self.assertTrue(all(run.bold for run in committee.runs if run.text.strip()))
        self.assertTrue(any(run.italic and run.text == "决定" for run in _paragraph(result, "第一条").runs))

    def test_working_paper_uses_the_measured_pitch_and_is_justified(self):
        document = Document()
        for text in ("工作文件1.0", "委员会：联合国教科文组织", "议题：测试", "起草国：法兰西共和国", "1. 鼓励各国合作；"):
            document.add_paragraph(text).runs[0].font.size = Pt(12)
        result = _run("working-paper", _save(document))
        clause = _paragraph(result, "1.")
        # Handbook 页31–32: 15.5 pt between baselines, exact for plain text.
        self.assertAlmostEqual(clause.paragraph_format.line_spacing.pt, 15.5, places=2)
        self.assertEqual(clause.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertEqual(clause.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    def test_friendly_amendment_header_is_body_size_bold(self):
        document = Document()
        for text in ("友好修正案1.0", "起草国：法兰西共和国", "附议国：日本国", "对决议草案第一条："):
            document.add_paragraph(text)
        result = _run("friendly-amendment", _save(document))
        # Handbook 页52–53: the title reads "修正案", set at 12 pt and bold.
        title = _paragraph(result, "修正案1.0")
        self.assertTrue(all(run.bold and not run.italic and run.font.size.pt == 12 for run in title.runs if run.text))

    def test_scrambled_run_sizes_all_become_the_handbook_size(self):
        document = Document()
        texts = RESOLUTION_HEADER + ["回顾以往决议，"] + [f"第{n}条 决定采取第{n}项措施；" for n in "一二三四五"]
        for index, text in enumerate(texts):
            paragraph = document.add_paragraph(text)
            paragraph.runs[0].font.size = Pt((6, 26, 9, 28)[index % 4])
            mark = paragraph._p.get_or_add_pPr()
            rpr = OxmlElement("w:rPr")
            size = OxmlElement("w:sz")
            size.set(qn("w:val"), "28")
            rpr.append(size)
            mark.append(rpr)
        result = _run("draft-resolution", _save(document))
        sizes = {run.font.size.pt for paragraph in _output(result).paragraphs for run in paragraph.runs if run.text.strip()}
        self.assertEqual(sizes, {12})


if __name__ == "__main__":
    unittest.main()
