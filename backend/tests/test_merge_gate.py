"""Regressions required by the v1.6.3 acceptance gate agreed with Codex.

Numbering values, cross references, footnotes, signing lines, idempotence,
line pitch around objects, the strict content guard and the failure classes.
"""

import base64
import sys
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app import content_guard  # noqa: E402
from app.docx_view import visible_text  # noqa: E402
from app.main import app  # noqa: E402
from app.pipelines import PIPELINES  # noqa: E402

TEMPLATES = ROOT / "templates" / "pkunmun2026"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
HEADER = ["决议草案1.2", "联合国大会", "巴勒斯坦问题", "起草国：黎巴嫩共和国", "附议国：伊拉克王国", "联合国大会，", "回顾相关原则，"]


def _save(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _source(lines) -> Document:
    document = Document()
    for text in lines:
        document.add_paragraph(text)
    return document


def _run(document_type, content, **kwargs):
    return PIPELINES[document_type](TEMPLATES).run(content, **kwargs)


def _errors(result):
    return [item for item in result.validations if item.status == "error"]


def _texts(content: bytes) -> list[str]:
    return [visible_text(p) for p in Document(BytesIO(content)).paragraphs]


class NumberingValueTests(unittest.TestCase):
    def test_typed_article_gap_and_cross_reference_are_kept_and_reported(self):
        lines = HEADER + ["第一条 决定成立委员会；", "第二条 请委员会依照第五条提交报告；", "第五条 决定继续审议。"]
        result = _run("draft-resolution", _save(_source(lines)))
        self.assertEqual(_errors(result), [])
        texts = _texts(result.content)
        self.assertTrue(any(text.startswith("第五条 决定继续审议") for text in texts))
        self.assertTrue(any("依照第五条" in text for text in texts))
        self.assertFalse(any(text.startswith("第三条") for text in texts), "a gap must not be filled")
        self.assertTrue(any("编号不连续" in warning and "第五条" in warning for warning in result.model.warnings))

    def test_word_list_keeps_its_definition_start_and_override(self):
        document = _numbered_source()
        result = _run("working-paper", _save(document))
        self.assertEqual(_errors(result), [])
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            numbering = archive.read("word/numbering.xml").decode("utf-8")
            body = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:numId="77"', numbering)
        self.assertIn("<w:startOverride w:val=\"4\"/>", numbering)
        self.assertIn('<w:numId w:val="77"/>', body)


class MarkerOverrideTests(unittest.TestCase):
    def test_override_layer_marker_is_normalized_and_values_kept(self):
        document = _numbered_source(override_font="Arial", override_size="16", start_override="5")
        result = _run("working-paper", _save(document))
        self.assertEqual(_errors(result), [])
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            numbering = archive.read("word/numbering.xml").decode("utf-8")
        from lxml import etree

        root = etree.fromstring(numbering.encode("utf-8"))
        num = next(node for node in root.findall(qn("w:num")) if node.get(qn("w:numId")) == "77")
        override = num.find(qn("w:lvlOverride"))
        self.assertEqual(override.find(qn("w:startOverride")).get(qn("w:val")), "5")
        level = override.find(qn("w:lvl"))
        self.assertEqual(level.find(qn("w:numFmt")).get(qn("w:val")), "decimal")
        self.assertEqual(level.find(qn("w:lvlText")).get(qn("w:val")), "%1.")
        rpr = level.find(qn("w:rPr"))
        self.assertEqual(rpr.find(qn("w:sz")).get(qn("w:val")), "24")
        self.assertEqual(rpr.find(qn("w:rFonts")).get(qn("w:ascii")), "Times New Roman")
        self.assertEqual(rpr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "SimSun")
        base = next(node for node in root.findall(qn("w:abstractNum")) if node.get(qn("w:abstractNumId")) == "70")
        self.assertEqual(base.find(f"{qn('w:lvl')}/{qn('w:rPr')}/{qn('w:sz')}").get(qn("w:val")), "24")


def _numbered_source(override_font=None, override_size=None, start_override="4") -> Document:
    """A working paper whose Word list starts at 4 through an override."""

    document = _source(["工作文件1.0", "《联合国气候变化框架公约》缔约方大会", "气候议题", "起草国：法兰西共和国"])
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "70")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1.")):
        node = OxmlElement(tag)
        node.set(qn("w:val"), value)
        level.append(node)
    abstract.append(level)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "77")
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), "70")
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), start_override)
    override.append(start)
    if override_font or override_size:
        # An override level that prints the number at 8 pt in another font.
        level_override = OxmlElement("w:lvl")
        level_override.set(qn("w:ilvl"), "0")
        for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1.")):
            node = OxmlElement(tag)
            node.set(qn("w:val"), value)
            level_override.append(node)
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), override_font or "Arial")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), override_size or "16")
        rpr.extend([fonts, size])
        level_override.append(rpr)
        override.append(level_override)
    num.extend([reference, override])
    numbering.insert(0, abstract)
    numbering.append(num)
    for text in ("鼓励各国合作；", "呼吁提交报告。"):
        paragraph = document.add_paragraph(text)
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = 77
    return document


class NotesAndPitchTests(unittest.TestCase):
    def test_inline_picture_paragraph_uses_a_minimum_pitch(self):
        tmp = Path(__file__).resolve().parent / "_gate_tmp"
        tmp.mkdir(exist_ok=True)
        try:
            picture = tmp / "pixel.png"
            picture.write_bytes(PNG)
            document = _source(HEADER + ["第一条 决定附图如下："])
            document.add_paragraph().add_run().add_picture(str(picture), width=Inches(1))
            document.add_paragraph("第二条 决定继续审议。")
            result = _run("draft-resolution", _save(document))
        finally:
            for item in tmp.glob("*"):
                item.unlink()
            tmp.rmdir()
        self.assertEqual(_errors(result), [])
        output = Document(BytesIO(result.content))
        with_picture = next(p for p in output.paragraphs if p._p.xpath(".//w:drawing"))
        self.assertEqual(with_picture.paragraph_format.line_spacing_rule, WD_LINE_SPACING.AT_LEAST)
        text = next(p for p in output.paragraphs if visible_text(p).startswith("第二条"))
        self.assertEqual(text.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)

    def test_footnote_text_is_set_at_nine_point(self):
        content = _with_footnote(_save(_source(["Position Paper", "Committee: GA", "Topic: Drugs", "Country: Afghanistan", "Delegate: Li", "Soft drugs are debated."])))
        result = _run("position-paper", content)
        self.assertEqual(_errors(result), [])
        with zipfile.ZipFile(BytesIO(result.content)) as archive:
            notes = archive.read("word/footnotes.xml").decode("utf-8")
        self.assertIn('w:sz w:val="18"', notes)
        self.assertIn("Reference entry.", notes)


def _with_footnote(content: bytes) -> bytes:
    """Add word/footnotes.xml with one note referenced from the last paragraph."""

    source = zipfile.ZipFile(BytesIO(content))
    target = BytesIO()
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    notes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{w}">'
        '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:sz w:val="30"/></w:rPr><w:t>Reference entry.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdNotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            elif item.filename == "word/document.xml":
                data = data.replace(b"<w:t>Soft drugs are debated.</w:t></w:r>", b'<w:t>Soft drugs are debated.</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r>')
            out.writestr(item, data)
        out.writestr("word/footnotes.xml", notes)
    return target.getvalue()


class SignatureLineTests(unittest.TestCase):
    def test_lists_break_between_whole_names_and_stay_stable(self):
        long_name = "大不列颠及北爱尔兰联合王国大不列颠及北爱尔兰联合王国大不列颠"  # longer than a line
        lines = ["决议草案1.2", "联合国大会", "巴勒斯坦问题", "起草国：黎巴嫩共和国、" + long_name + "、伊朗国",
                 "附议国：阿富汗王国、埃及王国、埃塞俄比亚帝国、波兰共和国、南斯拉夫社会主义联邦共和国、沙特阿拉伯国",
                 "联合国大会，", "回顾相关原则，", "第一条 决定继续审议。"]
        first = _run("draft-resolution", _save(_source(lines)))
        self.assertEqual(_errors(first), [])
        texts = _texts(first.content)
        self.assertTrue(any(long_name in text for text in texts), "a long name is never cut")
        joined = "".join(text for text in texts if "国" in text and "决议" not in text)
        for name in ("黎巴嫩共和国", "伊朗国", "波兰共和国", "沙特阿拉伯国"):
            self.assertEqual(joined.count(name), 1)
        second = _run("draft-resolution", first.content)
        third = _run("draft-resolution", second.content)
        self.assertEqual(_texts(second.content), _texts(first.content))
        self.assertEqual(_texts(third.content), _texts(first.content), "blank lines must not accumulate")


class ContentGuardTests(unittest.TestCase):
    def _paragraph(self, xml_body: str):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph._p.append(_fragment(xml_body))
        return document, paragraph

    def test_markup_like_text_is_not_structure(self):
        _, paragraph = self._paragraph('<w:r><w:t>&lt;w:hyperlink&gt;</w:t></w:r>')
        tokens = content_guard.signature(paragraph._p)
        self.assertTrue(all(token[0] == "t" for token in tokens))

    def test_field_edit_must_produce_the_authorized_value(self):
        document = _source(["委员会：旧名称"])
        before = content_guard.Snapshot.take(document)
        paragraph = document.paragraphs[0]
        paragraph.runs[0].text = "委员会：别的名称"
        edit_log = {paragraph._p: content_guard.Edit("field", "committee", "委员会：新名称")}
        problems = content_guard.verify_format(before, document, edit_log, allowed_titles=(), labels=("委员会",))
        self.assertTrue(problems)

    def test_new_country_line_may_hold_only_the_expected_names(self):
        document = _source(["起草国：甲国、乙国"])
        before = content_guard.Snapshot.take(document)
        extra = OxmlElement("w:p")
        extra.append(_fragment('<w:r><w:t>丙国</w:t></w:r>'))
        document.paragraphs[0]._p.addnext(extra)
        edit_log = {
            document.paragraphs[0]._p: content_guard.Edit("countries", "sponsors", "起草国：甲国、乙国"),
            extra: content_guard.Edit("countries", "sponsors", "起草国：甲国、乙国", created=True),
        }
        problems = content_guard.verify_format(before, document, edit_log, allowed_titles=(), labels=())
        self.assertTrue(any("预期" in problem for problem in problems))

    def test_relationship_target_change_is_detected(self):
        content = _save(_source(["text"]))
        changed = _rewrite_part(content, "word/_rels/document.xml.rels", lambda data: data.replace(b'Target="styles.xml"', b'Target="other.xml"'))
        self.assertTrue(content_guard.verify_package(content, changed))


class WrappedContentTests(unittest.TestCase):
    """Fields, links, revisions and tabs around the text that the pass edits."""

    def _with(self, lines, xml_body, after=()):
        document = _source(lines)
        paragraph = document.add_paragraph()
        for child in _fragments(xml_body):
            paragraph._p.append(child)
        for text in after:
            document.add_paragraph(text)
        return _save(document)

    def test_field_codes_survive_an_ending_edit(self):
        content = self._with(HEADER, '<w:r><w:t xml:space="preserve">第一条 参见 </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                             '<w:r><w:instrText xml:space="preserve"> REF _Ref1 \\h </w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t>。</w:t></w:r>')
        result = _run("draft-resolution", content)
        self.assertEqual(_errors(result), [])
        instr = Document(BytesIO(result.content)).element.body.findall(".//" + qn("w:instrText"))
        self.assertEqual([node.text for node in instr], [" REF _Ref1 \\h "])

    def test_dropping_a_label_inside_a_hyperlink_keeps_the_link(self):
        document = _source(["决议草案1.0"])
        paragraph = document.add_paragraph()
        r_id = document.part.relate_to("https://example.org/a", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link = _fragment('<w:hyperlink><w:r><w:t>委员会：联合国大会</w:t></w:r></w:hyperlink>')
        link.set(qn("r:id"), r_id)
        paragraph._p.append(link)
        for text in ["议题：测试", "起草国：日本国", "联合国大会，", "回顾以往决议，", "第一条 决定设立观察团。"]:
            document.add_paragraph(text)
        result = _run("draft-resolution", _save(document))
        self.assertEqual(_errors(result), [])
        output = Document(BytesIO(result.content))
        links = output.element.body.findall(".//" + qn("w:hyperlink"))
        self.assertEqual(len(links), 1)
        self.assertEqual(output.part.rels[links[0].get(qn("r:id"))].target_ref, "https://example.org/a")

    def test_a_country_list_holding_a_tracked_insertion_is_kept_whole(self):
        content = self._with(["决议草案1.0", "联合国大会", "议题"],
                             '<w:r><w:t>附议国：</w:t></w:r><w:ins w:id="1" w:author="A"><w:r><w:t>加拿大、澳大利亚、</w:t></w:r></w:ins>',
                             ["乍得共和国", "联合国大会，", "回顾以往决议，", "第一条 决定设立观察团。"])
        result = _run("draft-resolution", content)
        self.assertEqual(_errors(result), [])
        joined = "\n".join(_texts(result.content))
        self.assertIn("加拿大、澳大利亚、", joined)
        self.assertIn("乍得共和国", joined)
        self.assertEqual(len(Document(BytesIO(result.content)).element.body.findall(".//" + qn("w:ins"))), 1)

    def test_a_typed_number_after_a_tab_loses_no_character(self):
        content = self._with(["工作文件1.6", "《联合国气候变化框架公约》缔约方大会", "气候议题", "起草国：法兰西共和国"],
                             '<w:r><w:t>6.</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>认为教育主权不得妥协；</w:t></w:r>', ["1.鼓励各国合作；"])
        result = _run("working-paper", content)
        self.assertEqual(_errors(result), [])
        joined = "\n".join(_texts(result.content))
        self.assertIn("认为教育主权不得妥协", joined)
        self.assertIn("鼓励各国合作", joined)

    def _tracked_title(self) -> bytes:
        document = Document()
        paragraph = document.add_paragraph()
        for child in _fragments('<w:ins w:id="100" w:author="Reviewer" w:date="2026-01-01T00:00:00Z"><w:r><w:t>非友好修正案1.3.2</w:t></w:r></w:ins>'):
            paragraph._p.append(child)
        for text in ["起草国：日本国", "在第一条后增加：“决定继续审议。”"]:
            document.add_paragraph(text)
        return _save(document)

    def test_text_inside_a_tracked_revision_is_never_rewritten(self):
        result = _run("unfriendly-amendment", self._tracked_title())
        self.assertEqual(_errors(result), [])
        inserted = Document(BytesIO(result.content)).element.body.findall(".//" + qn("w:ins"))
        self.assertEqual("".join(node.text for node in inserted[0].iter(qn("w:t"))), "非友好修正案1.3.2")
        self.assertTrue(any("第 1 段" in (item.detail or "") and "修订" in (item.detail or "") for item in result.validations))

    def test_step_three_title_change_on_a_revision_is_refused_by_paragraph(self):
        from app.errors import ProtectedContentError

        with self.assertRaises(ProtectedContentError) as caught:
            _run("unfriendly-amendment", self._tracked_title(), overrides={"title": "非友好修正案2.1"})
        self.assertIn("第 1 段", str(caught.exception))
        self.assertIn("修订", str(caught.exception))


def _fragments(xml_body: str):
    from lxml import etree

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    return list(etree.fromstring(f'<w:wrap xmlns:w="{w}">{xml_body}</w:wrap>'))


def _fragment(xml_body: str):
    from lxml import etree

    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    return etree.fromstring(f'<w:wrap xmlns:w="{w}">{xml_body}</w:wrap>')[0]


def _rewrite_part(content: bytes, name: str, change) -> bytes:
    source = zipfile.ZipFile(BytesIO(content))
    target = BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for item in source.infolist():
            data = source.read(item.filename)
            out.writestr(item, change(data) if item.filename == name else data)
    return target.getvalue()


class FailureClassTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def _post(self, content: bytes, **fields):
        return self.client.post(
            "/api/format/draft-resolution",
            files={"file": ("a.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data=fields,
        )

    def test_damaged_file_is_rejected_as_unreadable(self):
        response = self._post(b"PK\x03\x04 not a docx")
        self.assertEqual(response.status_code, 422)

    def test_program_defect_is_reported_as_such(self):
        from app import pipelines

        original = pipelines.BasePipeline.run
        pipelines.BasePipeline.run = lambda *args, **kwargs: (_ for _ in ()).throw(KeyError("boom"))
        try:
            response = self._post(_save(_source(HEADER + ["第一条 决定。"])))
        finally:
            pipelines.BasePipeline.run = original
        self.assertEqual(response.status_code, 500)
        self.assertIn("程序内部错误", response.json()["detail"])


if __name__ == "__main__":
    unittest.main()
