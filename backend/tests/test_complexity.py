"""Formatting must stay linear in the number of paragraphs.

``Document.paragraphs`` rebuilds every ``Paragraph`` object on each access;
indexing it inside a loop once made a 2 000 paragraph resolution take ~29 s.
Wall-clock limits are flaky, so this counts the rebuilds instead: the count
must not grow with the document.
"""

import copy
import sys
import unittest
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

import docx.blkcntnr as blkcntnr  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.pipelines import PIPELINES  # noqa: E402

TEMPLATES = ROOT / "templates" / "pkunmun2026"
SAMPLES = ROOT / "examples" / "acceptance-inputs"
CASES = {
    "position-paper": "01_中文立场文件.docx",
    "working-paper": "03_中文工作文件.docx",
    "draft-directive": "05_中文指令草案.docx",
    "draft-resolution": "07_中文决议草案.docx",
    "friendly-amendment": "09_中文友好修正案.docx",
    "unfriendly-amendment": "10_中文非友好修正案.docx",
}


def repeated_body(content: bytes, times: int) -> bytes:
    document = Document(BytesIO(content))
    body = document.element.body
    blocks = [child for child in body if child.tag != qn("w:sectPr")]
    section = body.find(qn("w:sectPr"))
    for _ in range(times - 1):
        for block in blocks:
            section.addprevious(copy.deepcopy(block))
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


class ParagraphRebuildTests(unittest.TestCase):
    def setUp(self):
        self.original = blkcntnr.BlockItemContainer.paragraphs
        self.rebuilds = 0

        def counting(container):
            self.rebuilds += 1
            return self.original.fget(container)

        blkcntnr.BlockItemContainer.paragraphs = property(counting)

    def tearDown(self):
        blkcntnr.BlockItemContainer.paragraphs = self.original

    def rebuilds_for(self, document_type: str, content: bytes) -> int:
        self.rebuilds = 0
        PIPELINES[document_type](TEMPLATES).run(content)
        return self.rebuilds

    def test_paragraph_list_rebuilds_do_not_grow_with_document_size(self):
        for document_type, sample in CASES.items():
            with self.subTest(document_type=document_type):
                content = (SAMPLES / sample).read_bytes()
                small = self.rebuilds_for(document_type, content)
                large = self.rebuilds_for(document_type, repeated_body(content, 25))
                self.assertEqual(small, large)
                self.assertLessEqual(large, 6)


if __name__ == "__main__":
    unittest.main()
