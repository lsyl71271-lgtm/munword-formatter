"""Independent readback of edge cases found after the v1.8.4 baseline."""
import sys
import struct
import unittest
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'backend'))
from app.pipelines import PIPELINES
from app.treaty import is_treaty
from app.numbering import marker_of, valid_marker_change
from app.docx_package import validate_docx_package
from app.errors import InvalidDocxError

# The handbook's six types; diplomatic agreements and joint statements have their own tests.
HANDBOOK_TYPES = [kind for kind in PIPELINES if not is_treaty(kind)]

TITLES = ['立场文件', '工作文件', '指令草案', '决议草案', '友好修正案', '非友好修正案']

def source(kind):
    document = Document()
    for text in [TITLES[HANDBOOK_TYPES.index(kind)], '委员会：联合国大会', '议题：合作', '国家：中国', '代表：测试', '起草国：中国', '附议国：法国']:
        document.add_paragraph(text)
    return document

def save(document):
    data = BytesIO()
    document.save(data)
    return data.getvalue()

class EdgeCaseTests(unittest.TestCase):
    def format(self, data, kind):
        result = PIPELINES[kind](ROOT / 'templates/pkunmun2026').run(data)
        self.assertEqual([(item.code, item.detail) for item in result.validations if item.status == 'error'], [])
        return result.content, Document(BytesIO(result.content))

    def test_decimal_quantities_and_compound_outlines_are_preserved(self):
        lines = ['1.5 亿美元用于合作。', '1.1 资金安排', '1.2.3 项目说明', '2．5 吨物资。']
        for kind in HANDBOOK_TYPES:
            with self.subTest(kind=kind):
                document = source(kind)
                for text in [*lines, '3. 要求落实。']:
                    document.add_paragraph(text)
                data, out = self.format(save(document), kind)
                _, again = self.format(data, kind)
                self.assertTrue(set(lines).issubset({p.text for p in out.paragraphs}))
                self.assertEqual([p.text for p in out.paragraphs], [p.text for p in again.paragraphs])
        self.assertFalse(valid_marker_change('1.5 亿美元。', '第一条5 亿美元。'))

    def test_blank_looking_range_markers_survive_and_export(self):
        for tag in ['bookmarkEnd', 'commentRangeStart', 'commentRangeEnd', 'permEnd', 'moveFromRangeEnd']:
            with self.subTest(tag=tag):
                document = source('draft-resolution')
                document.add_paragraph('第一条 要求合作。')
                element = OxmlElement('w:' + tag)
                element.set(qn('w:id'), '7')
                document.add_paragraph()._p.append(element)
                data, out = self.format(save(document), 'draft-resolution')
                _, again = self.format(data, 'draft-resolution')
                for result in (out, again):
                    self.assertEqual(len(list(result.element.body.iter(qn('w:' + tag)))), 1)

    def test_empty_native_items_and_marked_whitespace_are_preserved(self):
        document = source('working-paper')
        for text in ['提交计划。', '', '提交报告。']:
            paragraph = document.add_paragraph(text)
            num = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            num.get_or_add_numId().val = 5
            num.get_or_add_ilvl().val = 0
        document.add_paragraph().add_run(' ').font.hidden = True
        data, out = self.format(save(document), 'working-paper')
        _, again = self.format(data, 'working-paper')
        for result in (out, again):
            self.assertEqual(len(list(result.element.body.iter(qn('w:numId')))), 3)
            self.assertEqual(len(list(result.element.body.iter(qn('w:vanish')))), 1)

    def test_partial_numbering_properties_keep_inherited_id(self):
        for mode in ['direct-level', 'derived-level', 'disabled', 'table']:
            with self.subTest(mode=mode):
                document = source('working-paper')
                paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if mode == 'table' else document.add_paragraph()
                paragraph.text = '正文列表内容。'
                paragraph.style = 'List Number'
                if mode == 'derived-level':
                    derived = document.styles.add_style('DerivedList', 1)
                    derived.base_style = document.styles['List Number']
                    derived.element.get_or_add_pPr().get_or_add_numPr().get_or_add_ilvl().val = 0
                    paragraph.style = derived
                else:
                    num = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    if mode == 'disabled':
                        num.get_or_add_numId().val = 0
                    else:
                        num.get_or_add_ilvl().val = 0
                data, out = self.format(save(document), 'working-paper')
                _, again = self.format(data, 'working-paper')
                for result in (out, again):
                    paragraph = next(p for p in result.element.body.iter(qn('w:p')) if '正文列表内容' in ''.join(p.itertext()))
                    identifier = paragraph.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
                    self.assertIsNotNone(identifier)
                    self.assertEqual(identifier.get(qn('w:val')), '0' if mode == 'disabled' else '5')

    def test_very_long_numeric_prefixes_are_bounded(self):
        self.assertIsNone(marker_of('9' * 5000 + '. text'))
        self.assertEqual(marker_of('0' * 5000 + '3. text')['value'], 3)

    def test_zip_local_header_metadata_must_match_central_directory(self):
        for offset in (14, 18, 22):
            with self.subTest(offset=offset):
                data = bytearray(save(source('working-paper')))
                struct.pack_into('<I', data, offset, struct.unpack_from('<I', data, offset)[0] + 1)
                with self.assertRaises(InvalidDocxError):
                    validate_docx_package(bytes(data))

    def test_localized_or_missing_normal_style_can_be_formatted(self):
        for mode in ('localized', 'missing'):
            with self.subTest(mode=mode):
                document = source('working-paper')
                document.add_paragraph('正文内容必须保留。')
                normal = document.styles['Normal'].element
                if mode == 'localized':
                    normal.set(qn('w:styleId'), 'a')
                    normal.find(qn('w:name')).set(qn('w:val'), '正文')
                else:
                    normal.getparent().remove(normal)
                data, out = self.format(save(document), 'working-paper')
                _, again = self.format(data, 'working-paper')
                self.assertEqual([p.text for p in out.paragraphs], [p.text for p in again.paragraphs])
                self.assertIn('正文内容必须保留。', [p.text for p in out.paragraphs])
