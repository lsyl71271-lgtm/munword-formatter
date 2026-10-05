"""Cached, table-aware views over a python-docx document.

``Document.paragraphs`` rebuilds the whole ``Paragraph`` list on every access.
The formatters index into that property inside loops, which made a single
format run cost O(n^2) (a 2 000 paragraph resolution took ~29 s, 73 % of it
inside ``blkcntnr.paragraphs``).  ``body_paragraphs`` returns the same list
from a one-slot cache that is invalidated whenever the body gains or loses a
child element, so call sites keep their original shape.

The cache is per thread: the API runs each request in a worker thread, and a
process-wide slot let one request read another request's paragraph list
between the identity check and the return.

``all_paragraphs`` additionally walks table cells (including nested tables),
which ``Document.paragraphs`` never returns.  Font normalization and the
content fingerprints use it so that text inside tables stops being invisible
to both the formatter and its own safety checks.
"""

from __future__ import annotations

import threading

from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run


_CACHE = threading.local()

# Non-text OOXML that must survive formatting, with the name shown to users
# when it disappears.  Losing any of these is content loss even when every
# visible character is preserved.
STRUCTURE_TAGS = {
    "w:drawing": "图片",               # inline and floating images
    "w:pict": "图形 / 文本框",          # legacy VML pictures and text boxes
    "w:object": "嵌入对象",             # embedded OLE objects
    "w:hyperlink": "超链接",            # hyperlinks (relationship-backed)
    "w:footnoteReference": "脚注引用",
    "w:endnoteReference": "尾注引用",
    "w:commentReference": "批注",
    "w:fldSimple": "域",               # field codes (page numbers, cross references)
    "w:instrText": "域代码",            # complex field instructions
    "w:sectPr": "分节符 / 页面设置",     # section breaks and per-section page geometry
    "w:tbl": "表格",
    "w:bookmarkStart": "书签",
    # Tracked changes hold text that ``paragraph.text`` does not return, so
    # the text fingerprints cannot see them disappear either.
    "w:ins": "修订（插入）",
    "w:del": "修订（删除）",
}
# The subset that can live inside a paragraph's content.  A paragraph whose
# visible text is empty but which carries one of these is not empty.
_INLINE_STRUCTURE = frozenset(qn(tag) for tag in STRUCTURE_TAGS if tag not in ("w:sectPr", "w:tbl"))


# Block-level wrappers Word treats as transparent: their paragraphs are part of
# the running text.  Tables and text boxes are not (shared with the browser
# engine's ``flowParagraphs``).
_FLOW_CONTAINERS = frozenset(qn(tag) for tag in ("w:sdt", "w:sdtContent", "w:customXml"))


def flow_paragraph_elements(body) -> list:
    """Paragraphs of the running text in reading order.

    Body paragraphs plus those inside body-level content controls and custom
    XML wrappers.  ``Document.paragraphs`` skipped a content control, so its
    clauses were never recognized, formatted or checked.
    """

    out = []
    for child in body:
        if child.tag == qn("w:p"):
            out.append(child)
        elif child.tag in _FLOW_CONTAINERS:
            out.extend(flow_paragraph_elements(child))
    return out


def _flow_size(element) -> tuple:
    """Cheap change detector: child counts along the wrapper chain only."""

    return (len(element), *(_flow_size(child) for child in element.iterchildren(*_FLOW_CONTAINERS)))


def body_paragraphs(document: DocumentObject) -> list[Paragraph]:
    """The running-text paragraphs, cached between edits (see the module docstring)."""

    body = document.element.body
    size = _flow_size(body)
    if getattr(_CACHE, "body", None) is body and _CACHE.size == size:
        return _CACHE.paragraphs
    paragraphs = [Paragraph(element, document._body) for element in flow_paragraph_elements(body)]
    _CACHE.body = body
    _CACHE.size = size
    _CACHE.paragraphs = paragraphs
    return paragraphs


def invalidate(document: DocumentObject | None = None) -> None:
    """Drop the cache after inserting or removing paragraphs, and between runs."""

    _CACHE.body = None
    _CACHE.size = -1
    _CACHE.paragraphs = []


def _inside(element, tag: str, stop) -> bool:
    ancestor = element.getparent()
    while ancestor is not None and ancestor is not stop:
        if ancestor.tag == tag:
            return True
        ancestor = ancestor.getparent()
    return False


def table_paragraphs(document: DocumentObject) -> list[Paragraph]:
    """Every paragraph inside a table cell, nested tables and tables in content
    controls included, each once (``row.cells`` repeated merged cells)."""

    body = document.element.body
    return [Paragraph(p, document._body) for p in body.iter(qn("w:p")) if _inside(p, qn("w:tc"), body)]


def all_paragraphs(document: DocumentObject) -> list[Paragraph]:
    """Every paragraph in document order: running text, table cells, text boxes.

    The same set the browser engine normalizes fonts and style emphasis on.
    """

    return [Paragraph(p, document._body) for p in document.element.body.iter(qn("w:p"))]


_HIDDEN_CONTAINERS = frozenset(qn(tag) for tag in ("w:del", "w:moveFrom", "w:txbxContent"))


def visible_runs(paragraph: Paragraph) -> list[Run]:
    """Every run whose text the reader sees, in document order.

    ``Paragraph.runs`` returns direct children only, so text inside a
    hyperlink, a tracked insertion or a content control used to escape font
    normalization and the emphasis checks.  Deleted revisions and text boxes
    (which hold paragraphs of their own) are left out.
    """

    return [Run(element, paragraph) for element in _visible_run_elements(paragraph._p)]


def _visible_run_elements(root) -> list:
    runs = []
    for element in root.iter(qn("w:r")):
        ancestor = element.getparent()
        if ancestor is root:
            runs.append(element)
            continue
        hidden = False
        while ancestor is not None and ancestor is not root:
            if ancestor.tag in _HIDDEN_CONTAINERS:
                hidden = True
                break
            ancestor = ancestor.getparent()
        if not hidden:
            runs.append(element)
    return runs


# Run children that carry visible text, as python-docx's Run.text reads them.
_RUN_TEXT_TAGS = frozenset(qn(tag) for tag in ("w:t", "w:tab", "w:br", "w:cr", "w:noBreakHyphen", "w:ptab"))


def visible_text(paragraph: Paragraph) -> str:
    """The paragraph's text as shown, including links and tracked insertions."""

    return "".join(
        str(child)
        for run in _visible_run_elements(paragraph._p)
        for child in run
        if child.tag in _RUN_TEXT_TAGS
    )


def active_num_id(paragraph: Paragraph) -> str | None:
    """The list a paragraph belongs to, or ``None``.

    ``numId`` 0 is Word's and WPS's explicit "numbering off" and must not be
    mistaken for list membership.
    """

    ppr = paragraph._p.pPr
    if ppr is None or ppr.numPr is None or ppr.numPr.numId is None:
        return None
    num_id = str(ppr.numPr.numId.val)
    return None if num_id == "0" else num_id


_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_INLINE_OBJECTS = frozenset(
    [qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:txbxContent"), f"{{{_MATH_NS}}}oMath", f"{{{_MATH_NS}}}oMathPara"]
)


def holds_inline_object(paragraph_element) -> bool:
    """True when a paragraph holds a picture, text box, embedded object or formula.

    Such content can be taller than a text line, so it must not sit in an
    exact line height (shared policy ``handbook.lineSpacing.rule``).
    """

    return any(node.tag in _INLINE_OBJECTS for node in paragraph_element.iter())


def carries_hidden_structure(paragraph: Paragraph) -> bool:
    """True when a paragraph holds content that its visible text does not show.

    That covers inline images, text boxes, fields, bookmarks, comments,
    tracked changes and the ``sectPr`` that defines the page geometry of
    everything before it.
    """

    element = paragraph._p
    if element.pPr is not None and element.pPr.find(qn("w:sectPr")) is not None:
        return True
    for node in element.iter():
        if node.tag in _INLINE_STRUCTURE:
            return True
        # A page or column break is the author's layout decision, not an empty line.
        if node.tag == qn("w:br") and node.get(qn("w:type")) in ("page", "column"):
            return True
    return False


_ON = ("1", "true", "on")


def style_index(document: DocumentObject) -> dict:
    return {style.get(qn("w:styleId")): style for style in document.styles.element.findall(qn("w:style"))}


def default_style_id(styles: dict, kind: str) -> str | None:
    """The style Word applies when nothing names one (``w:default``); Chinese Word calls the paragraph one "a"."""

    return next((style_id for style_id, style in styles.items()
                 if style.get(qn("w:type")) == kind and style.get(qn("w:default")) in _ON), None)


def style_chain(styles: dict, style_id):
    """``(id, element)`` for a style and the styles it is based on, nearest first."""

    seen = set()
    while style_id and style_id in styles and style_id not in seen:
        seen.add(style_id)
        yield style_id, styles[style_id]
        based = styles[style_id].find(qn("w:basedOn"))
        style_id = based.get(qn("w:val")) if based is not None else None


def own_runs(paragraph_element) -> list:
    """The paragraph's runs, without those of text-box paragraphs nested in it."""

    runs = []
    for run in paragraph_element.iter(qn("w:r")):
        parent = run.getparent()
        while parent is not None and parent.tag != qn("w:p"):
            parent = parent.getparent()
        if parent is paragraph_element:
            runs.append(run)
    return runs


def structure_signature(document: DocumentObject) -> dict[str, int]:
    """Count the non-text OOXML features that formatting must not remove."""

    root = document.element.body
    counts = {tag: len(root.findall(f".//{qn(tag)}")) for tag in STRUCTURE_TAGS}
    counts["w:t"] = len(root.findall(f".//{qn('w:t')}"))
    return counts


def structure_losses(before: dict[str, int], after: dict[str, int]) -> list[str]:
    """Human-readable list of structural features that disappeared."""

    losses: list[str] = []
    for tag, label in STRUCTURE_TAGS.items():
        lost = before.get(tag, 0) - after.get(tag, 0)
        if lost > 0:
            losses.append(f"{label} 减少 {lost} 处")
    return losses
